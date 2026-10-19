"""Синтетические Excel-шаблоны спецификации (реального образца в tests/fixtures/ нет)."""

import io
import zipfile

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

HEADERS = ["№", "Раздел", "Наименование", "Спецификация", "Ед.", "Кол-во", "Фото/ссылка товаров"]
WIDTHS = [6, 18, 30, 60, 8, 10, 24]
HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
DATA_FILL = PatternFill("solid", fgColor="EEF3FA")
THIN = Side(style="thin", color="999999")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def to_bytes(wb) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def style_header(ws, row=1, headers=HEADERS, start_col=1):
    for i, text in enumerate(headers):
        cell = ws.cell(row=row, column=start_col + i, value=text)
        cell.font = Font(name="Arial", bold=True, color="FFFFFF", size=11)
        cell.fill, cell.border = HEADER_FILL, BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for i, width in enumerate(WIDTHS):
        ws.column_dimensions[chr(ord("A") + start_col - 1 + i)].width = width


def style_data_row(ws, row, *, fill=DATA_FILL, height=30):
    for col in range(1, 8):
        cell = ws.cell(row=row, column=col)
        cell.font, cell.fill, cell.border = Font(name="Arial", size=10), fill, BORDER
        cell.alignment = Alignment(vertical="center")
    ws.cell(row=row, column=6).number_format = "0.00"
    ws.row_dimensions[row].height = height


def simple() -> bytes:
    """Шапка в первой строке, ниже — пустая размеченная строка-образец."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Спецификация"
    style_header(ws)
    style_data_row(ws, 2)
    return to_bytes(wb)


def with_existing_rows() -> bytes:
    """Второй лист; заголовок и титул выше шапки; раздел-строка и две строки данных; формула."""
    wb = Workbook()
    cover = wb.active
    cover.title = "Титул"
    cover["A1"] = "Приложение к договору"
    ws = wb.create_sheet("Спец")
    ws["A1"] = "Техническая спецификация"
    ws.merge_cells("A1:G1")
    style_header(ws, row=3)
    ws.merge_cells("A4:G4")
    ws["A4"] = "Раздел 1. Компьютеры"
    rows = [
        (1, "Раздел 1", "Ноутбук", "Процессор не менее 8 ядер", "шт", 2),
        (2, "Раздел 1", "Монитор", "Диагональ не менее 24 дюймов", "шт", 5),
    ]
    for i, values in enumerate(rows, start=5):
        for col, value in enumerate(values, start=1):
            ws.cell(row=i, column=col, value=value)
        style_data_row(ws, i, fill=PatternFill("solid", fgColor="FFF2CC"), height=42)
    ws["H20"] = "=SUM(F5:F6)"  # формула вне таблицы должна сохраниться
    return to_bytes(wb)


def merged_two_row_header() -> bytes:
    wb = Workbook()
    ws = wb.active
    style_header(ws, row=1)
    for col in range(1, 8):
        ws.merge_cells(start_row=1, end_row=2, start_column=col, end_column=col)
    style_data_row(ws, 3)
    return to_bytes(wb)


def header_variants() -> bytes:
    wb = Workbook()
    ws = wb.active
    style_header(
        ws,
        headers=[
            "No",
            "РАЗДЕЛ",
            "наименование товара",
            "СПЕЦИФИКАЦИЯ",
            "Ед. изм.",
            "Количество",
            "Фото / ссылка",
        ],
    )
    return to_bytes(wb)


def missing_columns() -> bytes:
    wb = Workbook()
    ws = wb.active
    style_header(ws, headers=["№", "Раздел", "Наименование", "Спецификация", "Ед."])
    return to_bytes(wb)


def merged_data_row() -> bytes:
    """Шапка из 8 колонок: «Спецификация» занимает D:E; строка-образец тоже объединяет D:E."""
    wb = Workbook()
    ws = wb.active
    headers = [
        "№",
        "Раздел",
        "Наименование",
        "Спецификация",
        "",
        "Ед.",
        "Кол-во",
        "Фото/ссылка товаров",
    ]
    for i, text in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=i, value=text or None)
        cell.font, cell.fill, cell.border = Font(bold=True), HEADER_FILL, BORDER
    ws.merge_cells("D1:E1")
    for col in range(1, 9):
        ws.cell(row=2, column=col).border = BORDER
    ws.merge_cells("D2:E2")
    return to_bytes(wb)


def with_footer(gap: int = 1) -> bytes:
    wb = Workbook()
    ws = wb.active
    style_header(ws)
    style_data_row(ws, 2)
    ws.cell(row=2 + gap, column=3, value="Итого")
    return to_bytes(wb)


def with_loss_objects() -> bytes:
    wb = Workbook()
    ws = wb.active
    style_header(ws)
    style_data_row(ws, 2)
    chart = BarChart()
    chart.add_data(Reference(ws, min_col=6, min_row=1, max_row=2), titles_from_data=True)
    ws.add_chart(chart, "J2")
    data = to_bytes(wb)
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as src, zipfile.ZipFile(out, "w") as dst:
        for item in src.infolist():
            dst.writestr(item, src.read(item.filename))
        dst.writestr("xl/media/image1.png", b"\x89PNG")
        dst.writestr("xl/vbaProject.bin", b"vba")
    return out.getvalue()
