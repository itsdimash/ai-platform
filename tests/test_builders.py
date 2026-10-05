import io

import pdfplumber
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation

from app.utils.docx_builder import build_document
from app.utils.pdf_builder import build_pdf
from app.utils.pptx_builder import build_presentation
from app.utils.xlsx_builder import build_spreadsheet


def test_pptx_reopens_with_expected_slides():
    data = build_presentation(
        "Стратегия 2026",
        "Kerneu",
        [
            {"title": "Цели", "content": ["Рост", "Маржа"]},
            {"title": "Риски", "content": ["Курс"]},
        ],
    )
    prs = Presentation(io.BytesIO(data))
    assert len(prs.slides) == 3
    assert prs.slides[0].shapes.title.text == "Стратегия 2026"
    assert prs.slides[1].shapes.title.text == "Цели"


def test_docx_reopens_with_headings_and_paragraphs():
    data = build_document(
        "Служебная записка",
        [{"heading": "Суть", "paragraphs": ["Первый абзац", "Второй абзац"]}],
    )
    doc = Document(io.BytesIO(data))
    texts = [p.text for p in doc.paragraphs]
    assert "Служебная записка" in texts and "Суть" in texts and "Второй абзац" in texts


def test_xlsx_reopens_and_numbers_stay_numeric():
    data = build_spreadsheet(
        [
            {
                "name": "Продажи",
                "headers": ["Товар", "Кол-во", "Цена"],
                "rows": [["Болт", "10", "2.5"]],
            }
        ]
    )
    ws = load_workbook(io.BytesIO(data))["Продажи"]
    assert [c.value for c in ws[1]] == ["Товар", "Кол-во", "Цена"]
    assert [c.value for c in ws[2]] == ["Болт", 10, 2.5]


def test_pdf_reopens_and_cyrillic_is_embedded():
    data = build_pdf(
        "Отчёт по складу",
        [
            {
                "heading": "Итоги",
                "paragraphs": ["Склад работает стабильно."],
                "bullets": ["Первый пункт", "Второй пункт"],
                "table": {"headers": ["Товар", "Остаток"], "rows": [["Болт", "10"]]},
            },
            {"paragraphs": ["Спецсимволы <b>&</b> не ломают разметку."]},
        ],
    )
    assert data.startswith(b"%PDF-")
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        text = "\n".join(page.extract_text() or "" for page in pdf.pages)
    assert "Отчёт по складу" in text
    assert "Склад работает стабильно." in text
    assert "Первый пункт" in text and "Болт" in text
    assert "<b>&</b>" in text
