import io

import pytest
from openpyxl import load_workbook

from app.utils.xlsx_template_filler import (
    SpecRow,
    XlsxTemplateError,
    fill_template,
    header_key,
    inspect_template,
)
from tests import xlsx_fixtures as fx


def row(name="Ноутбук", spec="Процессор не менее 8 ядер", qty=2, **kw):
    return SpecRow(name=name, specification=spec, unit="шт", quantity=qty, **kw)


def reopen(result):
    return load_workbook(io.BytesIO(result.data))


# --- заголовки и раскладка -----------------------------------------------------------------


@pytest.mark.parametrize(
    ("text", "key"),
    [
        ("№", "num"),
        ("No", "num"),
        ("№ п/п", "num"),
        ("п/п", "num"),
        ("Раздел", "section"),
        ("наименование товара", "name"),
        ("СПЕЦИФИКАЦИЯ", "spec"),
        ("Ед.", "unit"),
        ("Ед. изм.", "unit"),
        ("Кол-во", "qty"),
        ("Количество", "qty"),
        ("Фото/ссылка товаров", "link"),
        ("Фото / ссылка", "link"),
        ("Срок поставки", None),
        (None, None),
        (123, None),
    ],
)
def test_header_key(text, key):
    assert header_key(text) == key


def test_simple_layout():
    layout = inspect_template(fx.simple())
    assert (layout.sheet, layout.header_row, layout.header_bottom) == ("Спецификация", 1, 1)
    assert layout.missing == [] and layout.start_row == 2 and layout.next_number == 1
    assert layout.columns["link"] == 7 and layout.warnings == []


def test_header_not_in_first_row_second_sheet_and_existing_rows():
    layout = inspect_template(fx.with_existing_rows())
    assert layout.sheet == "Спец" and layout.header_row == 3
    assert layout.start_row == 7  # после раздела (строка 4) и двух строк данных (5–6)
    assert layout.next_number == 3 and layout.style_row == 6


def test_merged_two_row_header_starts_after_merge():
    layout = inspect_template(fx.merged_two_row_header())
    assert layout.header_row == 1 and layout.header_bottom == 2 and layout.start_row == 3


def test_header_variants_are_recognised():
    layout = inspect_template(fx.header_variants())
    assert layout.missing == [] and set(layout.columns) == {
        "num",
        "section",
        "name",
        "spec",
        "unit",
        "qty",
        "link",
    }


def test_missing_columns_are_reported_by_name():
    layout = inspect_template(fx.missing_columns())
    assert layout.missing == ["Кол-во", "Фото/ссылка товаров"]
    with pytest.raises(XlsxTemplateError) as exc:
        fill_template(fx.missing_columns(), [row()])
    assert "Кол-во" in str(exc.value) and "Фото/ссылка товаров" in str(exc.value)


def test_no_header_at_all():
    from openpyxl import Workbook

    wb = Workbook()
    wb.active["A1"] = "что-то"
    buf = io.BytesIO()
    wb.save(buf)
    with pytest.raises(XlsxTemplateError) as exc:
        fill_template(buf.getvalue(), [row()])
    assert "не найдена строка заголовков" in str(exc.value)


def test_not_an_xlsx():
    with pytest.raises(XlsxTemplateError):
        inspect_template(b"not a workbook")


def test_loss_warnings_for_charts_media_and_macros():
    warnings = inspect_template(fx.with_loss_objects()).warnings
    text = " ".join(warnings)
    for label in ("рисунки", "диаграммы", "изображения", "макросы"):
        assert label in text


# --- запись ---------------------------------------------------------------------------------


def test_fill_simple_keeps_header_widths_and_copies_row_style():
    result = fill_template(fx.simple(), [row(), row("Монитор", qty=5)], default_section="Раздел 1")
    ws = reopen(result)["Спецификация"]
    assert (result.first_row, result.last_row, result.added) == (2, 3, 2)
    assert [c.value for c in ws[2]] == [
        1,
        "Раздел 1",
        "Ноутбук",
        "Процессор не менее 8 ядер",
        "шт",
        2,
        None,
    ]
    assert next(c.value for c in ws[3]) == 2 and ws["C3"].value == "Монитор"
    # шапка и ширины не тронуты
    assert ws["A1"].value == "№" and ws["A1"].fill.fgColor.rgb.endswith("1F3A5F")
    assert ws["D1"].font.bold and ws.column_dimensions["D"].width == 60
    # стиль строки-образца (под шапкой) скопирован на все новые строки
    for r in (2, 3):
        assert (
            ws.cell(row=r, column=3).font.name == "Arial"
            and ws.cell(row=r, column=3).border.left.style == "thin"
        )
        assert ws.cell(row=r, column=6).number_format == "0.00"
    assert ws["D3"].alignment.wrap_text is True and ws.row_dimensions[3].height == 30


def test_fill_continues_after_existing_rows_numbering_and_style_of_last_row():
    result = fill_template(fx.with_existing_rows(), [row("Принтер", qty=1), row("Сканер", qty=3)])
    wb = reopen(result)
    ws = wb["Спец"]
    assert (result.first_row, result.last_row) == (7, 8)
    assert [ws.cell(row=7, column=c).value for c in (1, 3, 6)] == [3, "Принтер", 1]
    assert ws["A8"].value == 4
    assert ws["A7"].fill.fgColor.rgb.endswith("FFF2CC") and ws.row_dimensions[7].height == 42
    # существующее не тронуто, формула сохранена, титульный лист на месте
    assert ws["C5"].value == "Ноутбук" and ws["A4"].value == "Раздел 1. Компьютеры"
    assert ws["H20"].value == "=SUM(F5:F6)"
    assert wb["Титул"]["A1"].value == "Приложение к договору"
    assert "A4:G4" in {str(r) for r in ws.merged_cells.ranges}


def test_section_priority_row_over_parameter():
    result = fill_template(
        fx.simple(), [row(section="Раздел 2"), row()], default_section="Раздел 1"
    )
    ws = reopen(result)["Спецификация"]
    assert ws["B2"].value == "Раздел 2" and ws["B3"].value == "Раздел 1"
    ws2 = reopen(fill_template(fx.simple(), [row()]))["Спецификация"]
    assert ws2["B2"].value is None


def test_links_hyperlink_placeholders_and_formula_text():
    rows = [
        row("A", link_url="https://shop.example/p/1"),
        row("B", link_note="[требуется проверка]"),
        row("C", link_note="[требуется веб-поиск]"),
        row("D"),
        row("=cmd()", spec="=1+1"),
    ]
    ws = reopen(fill_template(fx.simple(), rows))["Спецификация"]
    link = ws["G2"]
    assert link.value == "Ссылка" and link.hyperlink.target == "https://shop.example/p/1"
    assert link.font.underline == "single" and link.font.color.rgb.endswith("0563C1")
    assert ws["G3"].value == "[требуется проверка]" and ws["G4"].value == "[требуется веб-поиск]"
    assert ws["G5"].value is None and ws["G5"].hyperlink is None
    assert ws["C6"].data_type == "s" and ws["D6"].data_type == "s"  # не формулы


def test_control_characters_are_stripped():
    ws = reopen(fill_template(fx.simple(), [row(name="Стол\x00\x0b")]))["Спецификация"]
    assert ws["C2"].value == "Стол"


def test_merged_data_row_style_is_replicated_and_written_to_anchor():
    result = fill_template(fx.merged_data_row(), [row(), row()])
    ws = reopen(result).active
    ranges = {str(r) for r in ws.merged_cells.ranges}
    assert {"D1:E1", "D2:E2", "D3:E3"} <= ranges
    assert ws["D3"].value == "Процессор не менее 8 ядер" and ws["F3"].value == "шт"


def test_footer_collision_is_an_error_not_silent_overwrite():
    data = fx.with_footer(gap=2)  # «Итого» в строке 4, а запись 5 строк займёт 2–6
    with pytest.raises(XlsxTemplateError) as exc:
        fill_template(data, [row() for _ in range(5)])
    assert "4" in str(exc.value) and "перезаписала бы" in str(exc.value)
    result = fill_template(data, [row()])  # одна строка помещается
    assert reopen(result).active["C4"].value == "Итого"


def test_fill_does_not_mutate_input_and_zero_rows_is_noop():
    data = fx.simple()
    result = fill_template(data, [])
    assert result.added == 0 and reopen(result).active.max_row <= 2
