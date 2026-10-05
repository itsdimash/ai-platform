"""Раунд 2 по презентациям: плейсхолдеры, слова не рвутся, диаграммы, титульный без картинки,
stat, вертикальный баланс, сравнение, footnote/section, чередование layout'ов."""

import io
import logging

import pytest
from pptx import Presentation
from pptx.enum.chart import XL_LEGEND_POSITION

from app.system_prompt import build_system_prompt
from app.utils.deck import images as deck_images
from app.utils.deck import layout as L
from app.utils.deck.audit import audit_pptx
from app.utils.deck.insights import chart_takeaways
from app.utils.deck.schema import PRESENTATION_TOOL, Chart, has_placeholder, normalize, split_points
from app.utils.pptx_builder import prepare_deck, render

EMU = 914400


def build(args, images=None):
    deck, slides = prepare_deck(args)
    result = render(deck, slides, images or {})
    return deck, slides, result, Presentation(io.BytesIO(result.data))


def deck_of(*slides, **extra):
    return {"title": "Тест", "theme": "graphite", "summary": "s", "slides": list(slides), **extra}


def sl(layout, title="Заголовок слайда", **kw):
    return {"layout": layout, "title": title, "notes": "Заметка докладчика.", **kw}


def shapes(slide, name):
    return [s for s in slide.shapes if s.name == name or s.name.startswith(name)]


# --- 1. плейсхолдеры -------------------------------------------------------------------------------


@pytest.mark.parametrize(
    "text",
    [
        "Выручка ____ млн",
        "Рост [...] процентов",
        "Срок: TBD",
        "Цена XXX",
        "нужно заполнить",
        "Lorem ipsum",
        "???",
    ],
)
def test_placeholder_patterns_detected(text):
    assert has_placeholder(text)


@pytest.mark.parametrize(
    "text", ["Заполнитель для пустот", "Рост 18%", "Выручка 124 млн ₸", "Plan [1]"]
)
def test_normal_text_is_not_a_placeholder(text):
    assert not has_placeholder(text)


def test_bullet_with_placeholder_is_dropped_whole_and_counted(caplog):
    caplog.set_level(logging.WARNING)
    deck = normalize(
        deck_of(
            sl("bullets", bullets=["Рост на ____ процентов", "Хороший пункт", "Бюджет TBD млн"])
        )
    )
    assert deck.slides[1].bullets == ["Хороший пункт"]  # пункт удалён целиком, не вырезан кусок
    assert deck.stats["placeholder_dropped"] == 2
    assert sum("placeholder_dropped" in r.message for r in caplog.records) == 2


def test_placeholders_in_cards_steps_table_chart_quote_notes_and_titles():
    deck = normalize(
        deck_of(
            sl(
                "cards",
                cards=[
                    {"title": "А", "text": "ok"},
                    {"title": "Б", "text": "[...]"},
                    {"title": "В", "text": "ok"},
                ],
            ),
            sl(
                "timeline",
                steps=[
                    {"title": "1", "text": "x"},
                    {"title": "2", "text": "____"},
                    {"title": "3", "text": "z"},
                ],
            ),
            sl(
                "table",
                table={"headers": ["a", "b"], "rows": [["1", "2"], ["TBD", "4"], ["5", "6"]]},
            ),
            sl(
                "chart",
                chart={
                    "type": "bar",
                    "categories": ["x", "XXX", "z"],
                    "series": [{"name": "s", "values": [1, 2, 3]}],
                },
            ),
            sl("quote", quote={"text": "Цитата ____", "author": "A"}),
            {**sl("bullets", bullets=["a"]), "notes": "Первое предложение. Второе: ____. Третье."},
            sl("bullets", title="Итог ____", bullets=["a"]),
        )
    )
    cards, steps, table, chart, quote, notes, titled = deck.slides[1:]
    assert [c[0] for c in cards.cards] == ["А", "В"]
    assert [t[0] for t in steps.steps] == ["1", "3"]
    assert table.table.rows == [["1", "2"], ["5", "6"]]
    assert chart.chart.categories == ["x", "z"] and chart.chart.series[0][1] == [1.0, 3.0]
    assert quote.layout == "bullets"  # цитата с плейсхолдером убрана -> слайд без quote
    assert notes.notes == "Первое предложение. Третье."  # заметки: удалено предложение
    assert titled.title == "a"  # заголовок с плейсхолдером заменён запасным
    assert deck.stats["placeholder_dropped"] >= 7


def test_stat_slide_with_fewer_than_two_values_after_drop_becomes_bullets(caplog):
    caplog.set_level(logging.WARNING)
    deck = normalize(
        deck_of(
            sl(
                "stat",
                stats=[{"value": "24%", "label": "маржа"}, {"value": "____", "label": "рост"}],
            ),
        )
    )
    s = deck.slides[1]
    assert s.layout == "bullets" and s.bullets == ["24% — маржа"] and not s.stats
    assert deck.stats["stat_demoted"] == 1
    assert any("stat_demoted" in r.message for r in caplog.records)


def test_stat_with_cards_is_demoted_to_cards():
    deck = normalize(
        deck_of(
            sl(
                "stat",
                stats=[
                    {"value": "Real-time", "label": "обновление"},
                    {"value": "24%", "label": "маржа"},
                ],
                cards=[{"title": "А", "text": "а"}, {"title": "Б", "text": "б"}],
            )
        )
    )
    assert deck.slides[1].layout == "cards"
    assert deck.stats["stat_value_rejected"] == 1 and deck.stats["stat_demoted"] == 1


def test_summary_and_deck_title_placeholders_are_cleaned():
    deck = normalize(
        {
            **deck_of(sl("bullets", bullets=["a"])),
            "title": "Отчёт ____",
            "summary": "Готово. Цена TBD.",
        }
    )
    assert deck.title == "Заголовок слайда" and deck.summary == "Готово."


def test_audit_flags_placeholder_in_a_ready_file():
    from pptx.util import Inches

    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text_frame.text = "Итог: ____"
    buf = io.BytesIO()
    prs.save(buf)
    assert [i.kind for i in audit_pptx(buf).issues if i.kind == "placeholder"] == ["placeholder"]


def test_prompts_forbid_placeholders_and_set_currency_and_image_bans():
    desc = PRESENTATION_TOOL["description"]
    for needle in ("____", "TBD", "₸", "never leave blanks".lower()):
        assert needle in desc or needle in desc.lower()
    from datetime import date

    prompt = build_system_prompt("pm", date(2026, 10, 5))
    for needle in (
        "«____»",
        "ориентир ~20%",
        "тенге (₸)",
        "05.10.2026",
        "логотипы и реальные бренды",
        "узнаваемые люди",
        "layout=image",
    ):
        assert needle in prompt
    props = PRESENTATION_TOOL["parameters"]["properties"]
    for desc in (
        props["image_prompt"]["description"],
        props["slides"]["items"]["properties"]["image_prompt"]["description"],
    ):
        assert "logos" in desc and "recognizable" in desc


# --- 2. слова не рвутся посреди слова --------------------------------------------------------------------


def test_fit_size_shrinks_for_the_longest_word_and_fails_if_it_never_fits():
    word = "Предпринимательство"
    big = L.fit_size([word], 4.0, 5.0, 26, 14)
    assert big is not None and L.longest_word_width([word], big) <= 4.0
    assert L.fit_size([word], 1.2, 5.0, 26, 14) is None  # даже 14 pt не помещает слово


def test_timeline_with_many_steps_or_wide_words_changes_layout():
    steps = [{"title": f"Этап {i}", "text": "Описание этапа"} for i in range(1, 6)]
    _, slides, _, _ = build(deck_of(sl("timeline", steps=steps)))
    assert slides[1].variant == "h"  # 5 коротких шагов помещаются в одну строку (метрики Carlito)
    six = [{"title": f"Этап {i}", "text": "Описание этапа"} for i in range(1, 7)]
    _, slides6, _, _ = build(deck_of(sl("timeline", steps=six)))
    assert slides6[1].variant in ("h2", "v")
    word = "Сверхконкурентоспособность"
    assert L.text_width(word, 16, True) > L.step_columns(4)[0][1]  # в колонку из четырёх не влезает
    wide = [
        {"title": word, "text": "x"},
        {"title": "Ответственность", "text": "y"},
        {"title": "Цели", "text": "z"},
        {"title": "Старт", "text": "w"},
    ]
    _, slides, result, _ = build(deck_of(sl("timeline", steps=wide)))
    assert slides[1].variant != "h"
    assert audit_pptx(io.BytesIO(result.data)).of("word_overflow") == []


def test_cards_with_wide_words_fall_back_to_numbered_rows():
    word = "Сверхконкурентоспособность" * 2
    cards = [("Раз", word)] * 3
    assert not L.cards_geometry(cards, "grid").ok  # в колонку из трёх слово не влезает
    geom = L.pick_cards_variant(cards, "grid")
    assert geom is not None and geom.variant == "numbered" and geom.ok
    args = deck_of(sl("cards", cards=[{"title": t, "text": x} for t, x in cards]))
    *_, result, _ = build(args)
    assert audit_pptx(io.BytesIO(result.data)).of("word_overflow") == []


def test_audit_word_overflow_checks_every_text_box_not_only_tables():
    from pptx.util import Inches, Pt

    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    tb = s.shapes.add_textbox(Inches(1), Inches(1), Inches(1.0), Inches(1))
    tb.text_frame.word_wrap = True
    run = tb.text_frame.paragraphs[0].add_run()
    run.text = "Предпринимательство"
    run.font.size = Pt(20)
    buf = io.BytesIO()
    prs.save(buf)
    found = audit_pptx(buf).of("word_overflow")
    assert found and "Предпринимательство" in found[0].detail


def test_long_words_in_every_layout_stay_inside_their_boxes():
    long = "Электрификация"
    args = deck_of(
        sl("bullets", bullets=[f"{long} и {long}ность", "Коротко"]),
        sl("cards", cards=[{"title": long, "text": long}] * 2),
        sl(
            "stat", stats=[{"value": "1 234 567", "label": long}, {"value": "42%", "label": "рост"}]
        ),
        sl("two_column", left=[long], right=[long]),
        sl("comparison", cards=[{"title": long, "text": long}, {"title": "Б", "text": "б"}]),
        sl("timeline", steps=[{"title": long, "text": long}] * 3),
        sl("table", table={"headers": [long, "b"], "rows": [[long, "1"]]}),
    )
    _, _, result, _ = build(args)
    report = audit_pptx(io.BytesIO(result.data))
    assert report.of("word_overflow") == [] and report.of("overflow") == [], [
        (i.kind, i.detail) for i in report.issues
    ]


# --- 3. диаграммы -----------------------------------------------------------------------------------------


def chart_slide(prs_slide):
    return next(s for s in prs_slide.shapes if s.has_chart)


def test_pie_shows_category_names_in_legend_on_the_right_and_percent_labels():
    chart = {
        "type": "pie",
        "unit": "млн ₸",
        "categories": ["Алматы", "Астана", "Шымкент"],
        "series": [{"name": "Выручка", "values": [50, 30, 20]}],
    }
    _, _, _, prs = build(deck_of(sl("chart", chart=chart)))
    gf = chart_slide(prs.slides[1])
    c = gf.chart
    assert c.has_legend and c.legend.position == XL_LEGEND_POSITION.RIGHT
    assert list(c.plots[0].categories) == ["Алматы", "Астана", "Шымкент"]
    assert c.plots[0].data_labels.show_percentage and c.has_title  # единица — в заголовке диаграммы
    assert "млн ₸" in c.chart_title.text_frame.text.replace("\u202f", " ")


def test_bar_chart_axis_labels_unit_and_single_series_has_no_legend():
    chart = {
        "type": "bar",
        "unit": "млн ₸",
        "categories": ["Июль", "Август", "Сентябрь"],
        "series": [{"name": "Факт", "values": [36, 41, 47]}],
    }
    _, _, _, prs = build(deck_of(sl("chart", chart=chart)))
    c = chart_slide(prs.slides[1]).chart
    assert not c.has_legend
    va = c.value_axis
    assert (va.minimum_scale, va.maximum_scale, va.major_unit) == (0, 50, 10)
    assert va.has_title and va.axis_title.text_frame.text.replace("\u202f", " ") == "млн ₸"
    dl = c.plots[0].data_labels
    assert (
        dl.position is not None and dl.number_format == "#,##0" and not dl.number_format_is_linked
    )


def test_multi_series_line_has_legend_and_nice_axis():
    chart = {
        "type": "line",
        "categories": ["a", "b", "c"],
        "series": [
            {"name": "S1", "values": [1.2, 2.5, 3.1]},
            {"name": "S2", "values": [0.8, 1.9, 2.2]},
        ],
    }
    _, _, _, prs = build(deck_of(sl("chart", chart=chart)))
    c = chart_slide(prs.slides[1]).chart
    assert c.has_legend
    assert c.value_axis.major_unit and c.value_axis.maximum_scale >= 3.1


@pytest.mark.parametrize(
    "lo,hi,zero,expect",
    [(36, 47, True, (0, 50, 10)), (0, 100, True, (0, 120, 20)), (118, 134, False, (115, 135, 5))],
)
def test_nice_axis(lo, hi, zero, expect):
    assert L.nice_axis(lo, hi, zero) == expect


def test_chart_is_about_62_percent_wide_with_computed_takeaways_on_the_right(caplog):
    caplog.set_level(logging.WARNING)
    chart = {
        "type": "bar",
        "unit": "млн ₸",
        "categories": ["Июль", "Август", "Сентябрь"],
        "series": [{"name": "Факт", "values": [36, 41, 47]}],
    }
    deck, _, _, prs = build(deck_of(sl("chart", chart=chart)))
    gf = chart_slide(prs.slides[1])
    assert abs(gf.width / EMU / L.CONTENT_W - 0.62) < 0.01
    bullets = next(s for s in prs.slides[1].shapes if s.name == "Bullets")
    assert bullets.left > gf.left + gf.width  # выводы — справа, не под диаграммой
    texts = [p.text.replace("\u202f", " ") for p in bullets.text_frame.paragraphs]
    assert texts == [
        "Максимум: Сентябрь — 47 млн ₸",
        "Минимум: Июль — 36 млн ₸",
        "От «Июль» к «Сентябрь»: +11 млн ₸ (+30,6%)",
    ]
    assert deck.stats["chart_takeaways_computed"] == 1


def test_model_bullets_are_kept_when_given():
    chart = {"type": "bar", "categories": ["a", "b"], "series": [{"name": "s", "values": [1, 2]}]}
    deck, slides, *_ = build(
        deck_of(sl("chart", chart=chart, bullets=["Свой вывод 1", "Свой вывод 2"]))
    )
    assert (
        slides[1].bullets == ["Свой вывод 1", "Свой вывод 2"]
        and deck.stats["chart_takeaways_computed"] == 0
    )


def test_pie_takeaways_and_english():
    assert chart_takeaways(Chart("pie", ["A", "B", "C"], [("S", [50, 30, 20])]), "Share") == [
        "Largest share: A — 50%",
        "Smallest share: C — 20%",
        "Top two combined — 80%",
    ]
    ru = chart_takeaways(
        Chart("pie", ["Алматы", "Астана", "Шымкент"], [("S", [50, 30, 20])]), "Доли"
    )
    assert ru[0] == "Крупнейшая доля: Алматы — 50%"


# --- 4. титульный без картинки ----------------------------------------------------------------------------


def test_title_without_image_has_no_lone_circle_and_has_layered_composition():
    _, _, result, prs = build(deck_of(sl("bullets", bullets=["a"])))
    title = prs.slides[0]
    assert not [s for s in title.shapes if "circle" in s.name.lower() or "oval" in s.name.lower()]
    slabs = shapes(title, "Panel slab")
    assert len(slabs) >= 3 and len({s.fill.fore_color.rgb for s in slabs}) == len(slabs)
    assert audit_pptx(io.BytesIO(result.data)).issues == []


async def test_failed_image_generation_uses_the_same_composition(monkeypatch):
    async def boom(*a, **k):
        raise RuntimeError("x")

    monkeypatch.setattr(deck_images, "generate_image_bytes", boom)
    deck, slides = prepare_deck(deck_of(sl("bullets", bullets=["a"]), image_prompt="A cat"))
    images, _ = await deck_images.generate_deck_images(deck, slides)
    assert images == {}
    prs = Presentation(io.BytesIO(render(deck, slides, images).data))
    assert shapes(prs.slides[0], "Panel slab")


# --- 5. stat -------------------------------------------------------------------------------------------------


def run_sizes(shape):
    return [r.font.size.pt for r in shape.text_frame.paragraphs[0].runs]


def test_stat_values_share_one_size_and_unit_is_a_smaller_run_on_the_same_line():
    stats = [
        {"value": "124 млн ₸", "label": "выручка"},
        {"value": "+18%", "label": "рост"},
        {"value": "7 дней", "label": "срок"},
    ]
    *_, prs = build(deck_of(sl("stat", stats=stats)))
    values = shapes(prs.slides[1], "Stat value")
    nums = {run_sizes(v)[0] for v in values}
    assert len(nums) == 1  # один кегль на слайде
    first = values[0].text_frame.paragraphs[0]
    assert [r.text for r in first.runs] == ["124", "\u202fмлн\u202f₸"] and len(
        values[0].text_frame.paragraphs
    ) == 1  # одна строка
    assert run_sizes(values[0])[1] < run_sizes(values[0])[0]
    assert [r.text for r in values[1].text_frame.paragraphs[0].runs] == ["+18", "%"]


def test_stat_size_is_the_smallest_fitted_across_values():
    stats = [
        {"value": "7", "label": "a"},
        {"value": "1 234 567", "label": "b"},
        {"value": "42%", "label": "c"},
        {"value": "88", "label": "d"},
    ]
    *_, prs = build(deck_of(sl("stat", stats=stats)))
    sizes = {run_sizes(v)[0] for v in shapes(prs.slides[1], "Stat value")}
    assert len(sizes) == 1 and sizes.pop() < 60


def test_more_than_four_stats_are_split_and_tool_says_numbers_only():
    stats = [{"value": str(i), "label": "x"} for i in range(1, 6)]
    _, slides, *_ = build(deck_of(sl("stat", stats=stats)))
    assert all(len(s.stats) <= 4 for s in slides if s.layout == "stat")
    stat_desc = PRESENTATION_TOOL["parameters"]["properties"]["slides"]["items"]["properties"][
        "stats"
    ]["description"]
    assert "2-4" in stat_desc and "8 chars" in stat_desc
    assert "never words" in PRESENTATION_TOOL["description"]


# --- 6. вертикальный баланс ---------------------------------------------------------------------------------------


def test_short_bullet_list_is_large_and_accent_bar_equals_text_height():
    *_, prs = build(
        deck_of(sl("bullets", bullets=["Первый пункт", "Второй пункт", "Третий пункт"]))
    )
    s = prs.slides[1]
    text = shapes(s, "Bullets")[-1]
    size = text.text_frame.paragraphs[0].runs[0].font.size.pt
    assert size == 28  # до 4 пунктов — 28 pt
    bar = next(x for x in s.shapes if x.name == "Bullets accent")
    assert abs(bar.height - text.height) <= 0.03 * EMU and abs(bar.top - text.top) < 0.01 * EMU


def test_cards_height_fits_content_between_1_6_and_3_2_and_block_is_centered():
    cards = [
        {"title": "Один", "text": "Коротко"},
        {"title": "Два", "text": "Тоже коротко"},
        {"title": "Три", "text": "Ещё"},
    ]
    *_, prs = build(deck_of(sl("cards", cards=cards)))
    boxes = shapes(prs.slides[1], "Card")
    boxes = [b for b in boxes if b.name == "Card"]
    assert all(1.6 - 0.01 <= b.height / EMU <= 3.2 + 0.01 for b in boxes)
    top, bottom = boxes[0].top / EMU, (boxes[0].top + boxes[0].height) / EMU
    assert abs((top - L.CONTENT_Y) - (L.CONTENT_BOTTOM - bottom)) < 0.05  # по центру области


def test_audit_warns_when_content_fills_less_than_55_percent():
    *_, result, _ = build(deck_of(sl("bullets", bullets=["Один"])))
    report = audit_pptx(io.BytesIO(result.data))
    assert [w.kind for w in report.warnings] == ["sparse"] and report.issues == []
    full = deck_of(
        sl("bullets", bullets=["Длинный пункт про склад и логистику по регионам страны"] * 5)
    )
    *_, result, _ = build(full)
    assert audit_pptx(io.BytesIO(result.data)).warnings == []


# --- 7. сравнение ------------------------------------------------------------------------------------------------------


def test_comparison_points_are_separate_bullets_and_comma_lists_are_split(caplog):
    caplog.set_level(logging.WARNING)
    cards = [
        {"title": "Консервативный", "text": "Рост 8%, без новых затрат, риски минимальны"},
        {"title": "Агрессивный", "text": "Рост 20%; вложения 15 млн ₸; нужны два менеджера"},
    ]
    deck, _, _, prs = build(deck_of(sl("comparison", cards=cards)))
    texts = shapes(prs.slides[1], "Panel text")
    assert [p.text for p in texts[0].text_frame.paragraphs] == [
        "Рост 8%",
        "Без новых затрат",
        "Риски минимальны",
    ]
    assert len(texts[1].text_frame.paragraphs) == 3 and deck.stats["comparison_split"] == 2


def test_comparison_accepts_left_right_arrays():
    args = deck_of(
        sl(
            "comparison",
            cards=[{"title": "До", "text": ""}, {"title": "После", "text": ""}],
            left=["Медленно", "Дорого"],
            right=["Быстро"],
        )
    )
    *_, prs = build(args)
    texts = shapes(prs.slides[1], "Panel text")
    assert [len(t.text_frame.paragraphs) for t in texts] == [2, 1]


def test_comparison_panels_have_equal_fitted_height_and_headers_clear_of_vs():
    cards = [
        {"title": "Левый вариант", "text": "Раз; два"},
        {"title": "Правый вариант", "text": "Раз; два; три; четыре"},
    ]
    *_, prs = build(deck_of(sl("comparison", cards=cards)))
    s = prs.slides[1]
    panels = [x for x in s.shapes if x.name == "Panel"]
    assert panels[0].height == panels[1].height and panels[0].height / EMU < L.CONTENT_H
    vs = next(x for x in s.shapes if x.name == "VS")
    heads = shapes(s, "Panel title")
    assert (heads[1].left - (vs.left + vs.width)) / EMU >= 0.45 - 1e-6
    assert (vs.left - (heads[0].left + heads[0].width)) / EMU >= 0.45 - 1e-6


def test_split_points():
    assert split_points("a; b") == ["A", "B"]
    assert split_points("Если клиент платит вовремя, мы снижаем цену, а потом считаем") == [
        "Если клиент платит вовремя, мы снижаем цену, а потом считаем"
    ]


# --- 9-10. footnote и section --------------------------------------------------------------------------------------------


def test_footnote_is_12pt_muted_above_the_footer():
    *_, prs = build(deck_of(sl("bullets", bullets=["a", "b"], footnote="Условные данные")))
    s = prs.slides[1]
    note = shapes(s, "Footnote")[0]
    footer = shapes(s, "Footer")[0]
    assert (
        note.text_frame.text == "Условные данные"
        and note.text_frame.paragraphs[0].runs[0].font.size.pt == 12
    )
    assert (note.top + note.height) <= footer.top


def test_data_note_moves_from_title_subtitle_to_footnote(caplog):
    args = deck_of(
        sl("bullets", bullets=["a"]),
        sl("chart", chart={"categories": ["a", "b"], "series": [{"name": "s", "values": [1, 2]}]}),
        subtitle="Коммерческий департамент · Условные данные",
    )
    deck = normalize(args)
    assert deck.slides[0].subtitle == "Коммерческий департамент"
    assert deck.slides[2].footnote == "Условные данные" and deck.stats["data_note_moved"] == 1


def filler(n):
    return [sl("bullets", f"Слайд {i}", bullets=["x"]) for i in range(n)]


def test_sections_only_in_decks_of_twelve_or_more_with_number():
    small = normalize(deck_of(*filler(3), sl("section", "Раздел", subtitle="s")))
    assert all(s.layout != "section" for s in small.slides) and small.stats["section_dropped"] == 1
    big_args = deck_of(
        *filler(5),
        sl("section", "Раздел А", subtitle="Подзаголовок"),
        *filler(4),
        sl("section", "Раздел Б"),
    )
    _, slides, _, prs = build(big_args)
    sections = [s for s in slides if s.layout == "section"]
    assert [s.number for s in sections] == [1, 2]
    idx = slides.index(sections[0])
    number = shapes(prs.slides[idx], "Section number")[0]
    assert (
        number.text_frame.text == "01"
        and number.text_frame.paragraphs[0].runs[0].font.size.pt >= 80
    )


# --- 11. картинки ----------------------------------------------------------------------------------------------------------


async def test_every_image_prompt_gets_the_no_text_no_logo_no_people_suffix(monkeypatch):
    seen = []

    async def fake(prompt, size=None, quality=None):
        seen.append(prompt)
        return b"x", "m"

    monkeypatch.setattr(deck_images, "generate_image_bytes", fake)
    deck, slides = prepare_deck(
        deck_of(sl("image", bullets=["a"], image_prompt="A cat on a sofa"), image_prompt="A cat")
    )
    await deck_images.generate_deck_images(deck, slides)
    assert seen and all(p.endswith(deck_images.PROMPT_RULES) for p in seen)
    assert (
        "No text" in deck_images.PROMPT_RULES
        and "logos" in deck_images.PROMPT_RULES
        and "recognizable" in deck_images.PROMPT_RULES
    )


# --- 12. разнообразие --------------------------------------------------------------------------------------------------------


def layouts(slides):
    return [s.layout for s in slides]


def test_labelled_bullets_repeating_a_layout_become_cards(caplog):
    caplog.set_level(logging.WARNING)
    labelled = [
        "Склад: единый учёт остатков в реальном времени",
        "Закупки: автоматические заявки поставщикам",
        "Продажи: заказы и счета в одном окне",
    ]
    deck, slides, *_ = build(
        deck_of(sl("bullets", bullets=["Один", "Два"]), sl("bullets", "Модули", bullets=labelled))
    )
    assert layouts(slides) == ["title", "bullets", "cards"]
    assert deck.stats["layout_repeat_fixed"] + deck.stats["sparse_layout_changed"] == 1
    assert slides[2].cards[0] == ("Склад", "единый учёт остатков в реальном времени")


def test_two_labelled_groups_become_comparison():
    items = [
        "Плюсы: быстро и дёшево работает",
        "Плюсы: просто внедрить у себя",
        "Минусы: нужен обучающий период",
        "Минусы: ограничена кастомизация",
    ]
    _, slides, *_ = build(
        deck_of(sl("bullets", bullets=["Один", "Два"]), sl("bullets", "Итог", bullets=items))
    )
    assert slides[2].layout == "comparison" and slides[2].left == [
        "быстро и дёшево работает",
        "просто внедрить у себя",
    ]


def test_repeat_is_kept_and_logged_when_meaning_would_change(caplog):
    caplog.set_level(logging.WARNING)
    deck, slides, *_ = build(
        deck_of(
            sl("bullets", bullets=["Один", "Два"]),
            sl("bullets", "Просто", bullets=["Без двоеточий здесь", "И тут тоже"]),
        )
    )
    assert (
        layouts(slides) == ["title", "bullets", "bullets"] and deck.stats["layout_repeat_kept"] == 1
    )
    assert any("layout_repeat_kept" in r.message for r in caplog.records)


def test_cards_alternate_between_grid_and_numbered():
    def cards_slide(t):
        return sl(
            "cards",
            t,
            cards=[
                {"title": "Раз", "text": "Первый блок"},
                {"title": "Два", "text": "Второй блок"},
                {"title": "Три", "text": "Третий блок"},
            ],
        )

    _, slides, _, prs = build(
        deck_of(
            cards_slide("А"),
            sl("bullets", bullets=["x"]),
            cards_slide("Б"),
            sl("bullets", "В", bullets=["y"]),
            cards_slide("Г"),
        )
    )
    variants = [s.variant for s in slides if s.layout == "cards"]
    assert variants == ["grid", "numbered", "grid"]
    numbered_idx = next(i for i, s in enumerate(slides) if s.variant == "numbered")
    assert shapes(prs.slides[numbered_idx], "Card number 1")


def test_tool_description_hints_alternation_and_no_blanks():
    d = PRESENTATION_TOOL["description"]
    assert "never the same layout twice" in d and "12+" in d and "footnote" in d


# --- JSON-строки вместо массивов (реальный случай Claude) -------------------------------------------------------


def test_json_string_arrays_are_parsed_not_lost():
    import json as _json

    slides = _json.dumps(
        [
            {"layout": "bullets", "title": "A", "bullets": ["x", "y"]},
            {
                "layout": "cards",
                "title": "B",
                "cards": _json.dumps([{"title": "1", "text": "a"}, {"title": "2", "text": "b"}]),
            },
        ]
    )
    deck = normalize({"title": "T", "theme": "graphite", "summary": "s", "slides": slides})
    assert [s.layout for s in deck.slides] == ["title", "bullets", "cards"]
    assert deck.slides[2].cards == [("1", "a"), ("2", "b")]
    assert deck.stats["json_string_parsed"] == 2 and deck.stats["slides_not_list"] == 0


def test_broken_json_slides_string_is_salvaged_line_by_line():
    # реальный случай: у слайда с таблицей не закрыт объект table
    good = '{"layout":"bullets","title":"A","bullets":["x"]}'
    broken = (
        '{"layout":"table","title":"B","table":{"headers":["a","b"],"rows":[["1","2"]],"notes":"n"}'
    )
    lost = '{"layout": "bullets", "title": '
    deck = normalize(
        {
            "title": "T",
            "theme": "graphite",
            "summary": "s",
            "slides": f"[\n{good},\n{broken},\n{lost}\n]",
        }
    )
    assert [s.layout for s in deck.slides] == ["title", "bullets", "table"]
    assert deck.slides[2].table.rows == [["1", "2"]]
    assert deck.stats["json_slide_repaired"] == 1 and deck.stats["json_slide_lost"] == 1


def test_stat_value_split_keeps_ranges_together():
    from app.utils.deck.render import DeckRenderer

    split = DeckRenderer._split_value
    assert split("12–16 ч") == ("12–16", "ч")
    assert split("124 млн ₸") == ("124", "млн ₸")
    assert split("+18%") == ("+18", "%")
    assert split("1 234 567") == ("1 234 567", "")


def test_closing_block_is_vertically_centered_in_its_area():
    *_, prs = build(deck_of(sl("closing", bullets=["Шаг 1", "Шаг 2", "Шаг 3"])))
    block = shapes(prs.slides[1], "Bullets")[0]
    assert block.top / EMU > 2.9 + 0.2  # не прижат к верху области


def test_number_and_unit_stay_together_with_a_non_breaking_space():
    deck = normalize(deck_of(sl("bullets", bullets=["Выручка 152 млн ₸ за квартал", "Рост 12 %"])))
    assert deck.slides[1].bullets == ["Выручка 152\u202fмлн\u202f₸ за квартал", "Рост 12\u202f%"]
    # оценка переноса считает «152 млн ₸» одним словом — как настоящий рендерер
    lines = L.wrap_lines("Выручка 152\u202fмлн\u202f₸ за квартал", 24, 3.0)
    assert all("₸" not in line or "152" in line for line in lines)


def test_short_table_is_stretched_to_fill_the_area():
    args = deck_of(sl("table", table={"headers": ["А", "Б"], "rows": [["1", "2"], ["3", "4"]]}))
    *_, result, _ = build(args)
    report = audit_pptx(io.BytesIO(result.data))
    assert report.warnings == [] and report.issues == []


# --- раунд 3 ----------------------------------------------------------------------------------------------------------


from app.utils.deck.themes import get_palette, mix


def fact_plan_chart(fact=(380, 410, 450), plan=(390, 400, 420), **extra):
    return sl(
        "chart",
        "Выручка по месяцам",
        chart={
            "type": "bar",
            "unit": "млн ₸",
            "categories": ["Июль", "Август", "Сентябрь"],
            "series": [
                {"name": "Факт", "values": list(fact)},
                {"name": "План", "values": list(plan)},
            ],
        },
        **extra,
    )


def test_plan_fulfilment_claims_are_recomputed_from_the_chart(caplog):
    caplog.set_level(logging.WARNING)
    # 1240 / 1210 = 102,48% -> 102%
    deck = normalize(
        deck_of(
            fact_plan_chart(),
            sl(
                "stat",
                stats=[
                    {"value": "106%", "label": "Выполнение плана"},
                    {"value": "+12%", "label": "рост"},
                ],
            ),
            sl(
                "bullets",
                "План перевыполнен на 6%",
                bullets=["Квартал: план выполнен на 110%", "Рост 12%"],
            ),
            sl("closing", "Итоги", bullets=["Выполнение плана 106% за квартал"]),
        )
    )
    _, stat, bullets, closing = deck.slides[1:]
    assert (
        stat.stats[0] == ("102%", "Выполнение плана") and stat.stats[1][0] == "+12%"
    )  # прочее не трогаем
    assert bullets.title == "План перевыполнен на 2%"
    assert bullets.bullets == ["Квартал: план выполнен на 102%", "Рост 12%"]
    assert closing.bullets == ["Выполнение плана 102% за квартал"]
    assert deck.stats["numeric_mismatch"] == 4
    assert sum("numeric_mismatch" in r.message for r in caplog.records) == 4


def test_plan_claim_within_one_point_is_left_alone():
    deck = normalize(deck_of(fact_plan_chart(), sl("bullets", bullets=["План выполнен на 103%"])))
    assert (
        deck.slides[2].bullets == ["План выполнен на 103%"] and deck.stats["numeric_mismatch"] == 0
    )


def test_no_fact_plan_chart_means_no_rewriting():
    other = sl(
        "chart", chart={"categories": ["a", "b"], "series": [{"name": "Выручка", "values": [1, 2]}]}
    )
    deck = normalize(
        deck_of(
            other,
            sl(
                "stat",
                stats=[
                    {"value": "106%", "label": "Выполнение плана"},
                    {"value": "1%", "label": "x"},
                ],
            ),
        )
    )
    assert deck.slides[2].stats[0][0] == "106%" and deck.stats["numeric_mismatch"] == 0


def test_ambiguous_fact_plan_charts_are_skipped():
    deck = normalize(
        deck_of(
            fact_plan_chart(),
            fact_plan_chart(fact=(1, 1, 1), plan=(1, 1, 1)),
            sl("bullets", bullets=["План выполнен на 50%"]),
        )
    )
    assert deck.slides[3].bullets == ["План выполнен на 50%"]


def test_pie_percent_sum_is_normalized_and_logged(caplog):
    caplog.set_level(logging.WARNING)

    def pie(values, unit=""):
        return sl(
            "chart",
            chart={
                "type": "pie",
                "unit": unit,
                "categories": list("abcd"[: len(values)]),
                "series": [{"name": "S", "values": values}],
            },
        )

    deck = normalize(deck_of(pie([46, 24, 18, 17], "%"), pie([50, 30, 20]), pie([400, 300, 300])))
    bad, ok, big = (s.chart.series[0][1] for s in deck.slides[1:])
    assert sum(bad) == pytest.approx(100) and bad[0] == pytest.approx(43.8, abs=0.1)
    assert ok == [50.0, 30.0, 20.0] and big == [400.0, 300.0, 300.0]  # не проценты — не трогаем
    assert deck.stats["pie_sum_mismatch"] == 1 and any(
        "pie_sum_mismatch" in r.message for r in caplog.records
    )


def test_one_short_footnote_only_on_first_data_slide(caplog):
    caplog.set_level(logging.WARNING)
    deck = normalize(
        deck_of(
            sl(
                "bullets",
                bullets=["Один", "Два"],
                footnote="Условные данные; сумма факта 412 млн ₸",
            ),
            fact_plan_chart(footnote="Условные данные"),
            sl("table", table={"headers": ["a"], "rows": [["1"]]}, footnote="Условные данные"),
        )
    )
    notes = [(s.layout, s.footnote) for s in deck.slides if s.footnote]
    assert notes == [("chart", "Условные данные")]  # одна, короткая, на первом слайде с данными
    assert deck.stats["footnote_deduped"] == 2 and deck.stats["footnote_moved"] == 1


def test_instructions_are_stripped_from_slides_and_moved_to_the_summary(caplog):
    caplog.set_level(logging.WARNING)
    deck = normalize(
        {
            **deck_of(
                fact_plan_chart(footnote="Условные данные — заменить фактическими перед показом"),
                sl(
                    "bullets",
                    bullets=[
                        "Цифры в презентации — условные, ориентиры для замены фактом",
                        "Выручка выросла в каждом месяце квартала",
                    ],
                ),
            ),
            "summary": "Готово.",
        }
    )
    assert deck.slides[1].footnote == "Условные данные"
    assert deck.slides[2].bullets == [
        "Выручка выросла в каждом месяце квартала"
    ]  # пункт-инструкция удалён
    assert deck.stats["instruction_stripped"] == 2
    assert "замените фактическими" in deck.summary and any(
        "instruction_stripped" in r.message for r in caplog.records
    )


def test_notes_for_the_speaker_are_not_touched_by_instruction_stripping():
    deck = normalize(
        deck_of({**sl("bullets", bullets=["a"]), "notes": "Перед показом проверьте цифры."})
    )
    assert deck.slides[1].notes == "Перед показом проверьте цифры."


def test_stat_is_60pt_card_fits_content_is_centered_and_uses_real_minus_and_thin_space():
    stats = [
        {"value": "-30%", "label": "время отчётов"},
        {"value": "124 млн ₸", "label": "выручка"},
    ]
    *_, prs = build(deck_of(sl("stat", stats=stats)))
    s = prs.slides[1]
    values = shapes(s, "Stat value")
    assert values[0].text_frame.paragraphs[0].runs[0].text == "\u221230"  # настоящий минус
    assert [r.text for r in values[1].text_frame.paragraphs[0].runs] == ["124", "\u202fмлн\u202f₸"]
    assert values[0].text_frame.paragraphs[0].runs[0].font.size.pt == 60
    card = next(x for x in s.shapes if x.name == "Stat card")
    label = shapes(s, "Stat label")[0]
    assert (
        (label.top + label.height - card.top) / EMU
        < card.height / EMU
        <= (label.top + label.height - card.top) / EMU + 0.4
    )
    assert (
        abs((card.top / EMU - L.CONTENT_Y) - (L.CONTENT_BOTTOM - (card.top + card.height) / EMU))
        < 0.05
    )


def test_minus_is_real_only_before_a_number_not_in_ranges_or_words():
    deck = normalize(
        deck_of(
            sl(
                "bullets",
                bullets=["Падение -5% к плану", "Диапазон 10-15 дней", "COVID-19 и Wi-Fi"],
            )
        )
    )
    assert deck.slides[1].bullets == [
        "Падение \u22125% к плану",
        "Диапазон 10-15\u202fдней",
        "COVID-19 и Wi-Fi",
    ]


def test_image_less_title_uses_only_primary_tints_and_accent():
    p = get_palette("sunset")
    allowed = {p.primary, mix(p.primary, "FFFFFF", 0.45), mix(p.primary, "FFFFFF", 0.7), p.accent1}
    *_, prs = build({**deck_of(sl("bullets", bullets=["a"])), "theme": "sunset"})
    slabs = shapes(prs.slides[0], "Panel slab")
    assert len(slabs) == 4
    assert {str(x.fill.fore_color.rgb) for x in slabs} == allowed
    assert str(shapes(prs.slides[0], "Panel")[0].fill.fore_color.rgb) == p.dark  # без смеси с белым


def test_bullet_size_28_up_to_four_items_and_24_for_five():
    four = [
        "Пункт про склад и закупки раз",
        "Пункт про склад и закупки два",
        "Пункт про склад и закупки три",
        "Пункт про склад и закупки четыре",
    ]
    *_, prs = build(deck_of(sl("bullets", bullets=four)))
    assert shapes(prs.slides[1], "Bullets")[-1].text_frame.paragraphs[0].runs[0].font.size.pt == 28
    *_, prs = build(deck_of(sl("bullets", bullets=[*four, "Пункт про склад и закупки пять"])))
    assert shapes(prs.slides[1], "Bullets")[-1].text_frame.paragraphs[0].runs[0].font.size.pt == 24


def test_sparse_bullets_become_two_columns_or_cards_when_meaning_is_kept(caplog, monkeypatch):
    caplog.set_level(logging.WARNING)
    from app.utils.deck.layout import Planner

    monkeypatch.setattr(Planner, "SPARSE_BULLETS", 0.6)  # 4 коротких пункта — ~53% области
    short = ["Склад", "Закупки", "Продажи", "Финансы"]
    deck, slides, *_ = build(deck_of(sl("bullets", bullets=short)))
    assert slides[1].layout == "two_column" and slides[1].left == ["Склад", "Закупки"]
    assert deck.stats["sparse_layout_changed"] == 1
    labelled = ["Склад: единый учёт остатков", "Закупки: заявки поставщикам"]
    _, slides, *_ = build(deck_of(sl("bullets", bullets=labelled)))
    assert slides[1].layout == "cards"
    _, slides, *_ = build(
        deck_of(sl("bullets", bullets=["Один", "Два"]))
    )  # смысл не сохранить — остаётся
    assert slides[1].layout == "bullets"
    _, slides, *_ = build(deck_of(sl("image", bullets=["a", "b", "c", "d"], image_prompt="A cat")))
    assert slides[1].layout in ("image", "bullets")  # картинка — не трогаем


def test_text_is_measured_with_carlito_not_dejavu():
    from pathlib import Path

    assert (Path(L._FONT_DIR) / "Carlito-Regular.ttf").exists() and (
        Path(L._FONT_DIR) / "LICENSE_CARLITO.txt"
    ).exists()
    from PIL import ImageFont

    dejavu = ImageFont.truetype(str(Path(L._FONT_DIR) / "DejaVuSans.ttf"), 200)
    text = "Предпринимательство и конкурентоспособность рынка"
    old_estimate = dejavu.getlength(text) / 200 * 20 / 72 * 0.89  # прежняя оценка DejaVu x 0,89
    assert L.text_width(text, 20) < old_estimate  # новая оценка менее консервативна
    steps = [{"title": "Обследование", "text": "Разбираем процессы и цели клиента"}] * 5
    _, slides, *_ = build(deck_of(sl("timeline", steps=steps)))
    assert slides[1].variant == "h"


def test_footnote_sits_at_least_0_2_in_above_the_footer():
    *_, prs = build(deck_of(fact_plan_chart(footnote="Условные данные")))
    s = prs.slides[1]
    note, footer = shapes(s, "Footnote")[0], shapes(s, "Footer")[0]
    assert (footer.top - (note.top + note.height)) / EMU >= 0.2 - 1e-6


def test_three_cards_use_20pt_body_and_fit_height_to_content():
    cards = [
        {"title": t, "text": "Короткое описание блока в одну-две строки"}
        for t in ("Один", "Два", "Три")
    ]
    *_, prs = build(deck_of(sl("cards", cards=cards)))
    body = shapes(prs.slides[1], "Card text")[0].text_frame.paragraphs[0].runs[0].font.size.pt
    boxes = [b for b in shapes(prs.slides[1], "Card") if b.name == "Card"]
    assert body >= 20 and all(b.height / EMU <= 3.2 + 0.01 for b in boxes)
    assert (
        abs(
            (boxes[0].top / EMU - L.CONTENT_Y)
            - (L.CONTENT_BOTTOM - (boxes[0].top + boxes[0].height) / EMU)
        )
        < 0.05
    )


def test_shapes_have_no_theme_shadow_in_any_renderer():
    *_, prs = build(
        deck_of(sl("cards", cards=[{"title": "А", "text": "а"}, {"title": "Б", "text": "б"}]))
    )
    xml = prs.slides[1].shapes._spTree.xml
    assert 'effectRef idx="0"' in xml and 'effectRef idx="2"' not in xml


def test_flat_table_fields_build_a_table_and_legacy_nested_table_still_works():
    flat = normalize(
        deck_of(sl("table", table_headers=["А", "Б"], table_rows=[["1", "2"], ["3", "4"]]))
    )
    assert flat.slides[1].layout == "table" and flat.slides[1].table.rows == [
        ["1", "2"],
        ["3", "4"],
    ]
    nested = normalize(deck_of(sl("table", table={"headers": ["А"], "rows": [["1"]]})))
    assert nested.slides[1].table.headers == ["А"]
    as_strings = normalize(
        deck_of(sl("table", table_headers='["А","Б"]', table_rows='[["1","2"]]'))
    )
    assert as_strings.slides[1].table.rows == [["1", "2"]]
    props = PRESENTATION_TOOL["parameters"]["properties"]["slides"]["items"]["properties"]
    assert "table" not in props and {"table_headers", "table_rows"} <= set(props)
