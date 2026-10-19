"""Синтетические тендерные docx для тестов (реального образца в tests/fixtures/ нет)."""

import io

import docx

HEADER = (
    "Описание требуемых функциональных, технических, качественных, эксплуатационных "
    "и иных характеристик закупаемого товара"
)
LONG = (
    "Ноутбук - 2 шт. Процессор не менее 8 ядер, оперативная память не менее 16 ГБ, "
    "накопитель SSD не менее 512 ГБ, дисплей не менее 15 дюймов. "
    "Стол компьютерный - 5 шт. Столешница ЛДСП не менее 18 мм, ширина не менее 120 см."
)


def to_bytes(document) -> bytes:
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def layout_a(content: str = LONG, header: str = HEADER) -> bytes:
    """Заголовок в шапке колонки, содержимое — строкой ниже."""
    d = docx.Document()
    t = d.add_table(rows=2, cols=4)
    for cell, text in zip(t.rows[0].cells, ["№", "Наименование", header, "Срок"], strict=True):
        cell.text = text
    for cell, text in zip(
        t.rows[1].cells, ["1", "Закупка оборудования", content, "30 дней"], strict=True
    ):
        cell.text = text
    return to_bytes(d)


def layout_b(content: str = LONG) -> bytes:
    """Заголовок слева, содержимое справа; ниже другие пары «подпись — значение»."""
    d = docx.Document()
    t = d.add_table(rows=3, cols=2)
    rows = [
        ("Наименование закупки", "Оборудование"),
        (HEADER, content),
        ("Срок поставки", "30 дней"),
    ]
    for row, (left, right) in zip(t.rows, rows, strict=True):
        row.cells[0].text, row.cells[1].text = left, right
    return to_bytes(d)


def merged_horizontal() -> bytes:
    """Заголовок и содержимое объединены на две колонки (gridSpan)."""
    d = docx.Document()
    t = d.add_table(rows=2, cols=3)
    t.rows[0].cells[0].text = "№"
    t.rows[0].cells[1].merge(t.rows[0].cells[2]).text = HEADER
    t.rows[1].cells[0].text = "1"
    t.rows[1].cells[1].merge(t.rows[1].cells[2]).text = LONG
    return to_bytes(d)


def merged_vertical() -> bytes:
    """Содержимое занимает две строки (vMerge)."""
    d = docx.Document()
    t = d.add_table(rows=3, cols=2)
    t.rows[0].cells[0].text, t.rows[0].cells[1].text = "№", HEADER
    t.rows[1].cells[0].text = "1"
    t.rows[1].cells[1].merge(t.rows[2].cells[1]).text = LONG
    t.rows[2].cells[0].text = "2"
    return to_bytes(d)


def nested() -> bytes:
    """Нужная таблица лежит внутри ячейки другой таблицы."""
    d = docx.Document()
    outer = d.add_table(rows=1, cols=1)
    cell = outer.rows[0].cells[0]
    cell.text = "Приложение"
    inner = cell.add_table(rows=2, cols=2)
    inner.rows[0].cells[0].text, inner.rows[0].cells[1].text = "№", HEADER
    inner.rows[1].cells[0].text, inner.rows[1].cells[1].text = "1", LONG
    return to_bytes(d)


def two_tables() -> bytes:
    d = docx.Document()
    for text in ("Первое описание. Мышь - 3 шт.", "Второе описание. Клавиатура - 4 шт."):
        t = d.add_table(rows=2, cols=2)
        t.rows[0].cells[0].text, t.rows[0].cells[1].text = "№", HEADER
        t.rows[1].cells[0].text, t.rows[1].cells[1].text = "1", text
        d.add_paragraph("между таблицами")
    return to_bytes(d)


def no_cell() -> bytes:
    d = docx.Document()
    d.add_paragraph(HEADER)  # заголовок вне таблицы — не считается
    t = d.add_table(rows=2, cols=2)
    t.rows[0].cells[0].text, t.rows[0].cells[1].text = "№", "Срок"
    t.rows[1].cells[0].text, t.rows[1].cells[1].text = "1", "30 дней"
    return to_bytes(d)
