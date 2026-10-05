import asyncio
import io
import json
import logging

import pytest
from google.genai import types
from PIL import Image
from pptx import Presentation

from app.limits import DECK_MAX_IMAGES_HARD, MAX_SLIDES
from app.utils.deck import images as deck_images
from app.utils.deck import layout as L
from app.utils.deck.audit import audit_pptx
from app.utils.deck.render import prepare_picture
from app.utils.deck.schema import LAYOUTS, PRESENTATION_TOOL, normalize
from app.utils.deck.themes import (
    BODY_FLOOR,
    FONT,
    PALETTES,
    best_text_on,
    contrast,
    contrast_pairs,
    get_palette,
)
from app.utils.pptx_builder import build_presentation, prepare_deck, render
from tests.deck_fixtures import ALL_LAYOUTS, long_text_deck, one_layout
from tests.test_media import noisy_png


def build(args, images=None):
    deck, slides = prepare_deck(args)
    result = render(deck, slides, images or {})
    return deck, slides, result


def reopen(result):
    return Presentation(io.BytesIO(result.data))


# --- палитры -------------------------------------------------------------------------------


@pytest.mark.parametrize("name", list(PALETTES))
def test_every_palette_pair_passes_wcag_aa(name):
    for description, fg, bg, need in contrast_pairs(PALETTES[name]):
        assert contrast(fg, bg) >= need, f"{name}: {description} {contrast(fg, bg):.2f} < {need}"


def test_five_palettes_default_graphite_and_best_text_helper():
    assert set(PALETTES) == {"ocean", "forest", "sunset", "graphite", "kerneu"}
    assert get_palette(None).name == "graphite" == get_palette("nonexistent").name
    assert get_palette("KERNEU").dark == "0B2545"
    assert best_text_on("F2A900") == "1A1F24" and best_text_on("0B2545") == "FFFFFF"
    for p in PALETTES.values():
        assert len(set(p.chart_colors)) == len(p.chart_colors)  # серии диаграмм различимы


# --- каждый layout строится ---------------------------------------------------------------------


@pytest.mark.parametrize("layout", [layout for layout in LAYOUTS if layout != "title"])
def test_each_layout_builds_reopens_and_is_clean(layout):
    _, _, result = build(one_layout(layout))
    prs = reopen(result)
    assert len(prs.slides) >= 2 and result.slide_count == len(prs.slides)
    report = audit_pptx(io.BytesIO(result.data))
    assert report.issues == [], [(i.kind, i.detail) for i in report.issues]
    assert report.notes >= 1  # заметки докладчика на слайде
    last = prs.slides[-1]
    assert last.shapes.title.text.startswith(f"Слайд {layout}")  # настоящий title-placeholder


def test_title_layout_and_all_layouts_deck_is_clean_16_9():
    _, _, result = build(ALL_LAYOUTS)
    prs = reopen(result)
    assert (prs.slide_width, prs.slide_height) == (12192000, 6858000)  # 16:9
    assert len(prs.slides) == 12
    report = audit_pptx(io.BytesIO(result.data))
    assert report.issues == []
    assert (report.tables, report.charts) == (1, 1)
    assert report.notes == 12  # заметки на каждом слайде


def test_native_table_chart_and_notes_content():
    _, _, result = build(ALL_LAYOUTS)
    prs = reopen(result)
    table_slide = next(
        s
        for s in prs.slides
        if any(getattr(sh, "has_table", False) and sh.has_table for sh in s.shapes)
    )
    tbl = next(
        sh for sh in table_slide.shapes if getattr(sh, "has_table", False) and sh.has_table
    ).table
    assert [c.text.replace("\u202f", " ") for c in tbl.rows[0].cells] == [
        "Регион",
        "Выручка, млн ₸",
        "Доля",
        "Рост",
    ]
    assert len(tbl.rows) == 6
    chart_slide = next(
        s
        for s in prs.slides
        if any(getattr(sh, "has_chart", False) and sh.has_chart for sh in s.shapes)
    )
    chart = next(
        sh for sh in chart_slide.shapes if getattr(sh, "has_chart", False) and sh.has_chart
    ).chart
    assert [s.name for s in chart.plots[0].series] == ["Факт", "План"]
    assert list(chart.plots[0].series[0].values) == [36.0, 41.0, 47.0]
    assert prs.slides[1].notes_slide.notes_text_frame.text == "Главный слайд итогов."


@pytest.mark.parametrize("kind", ["bar", "line", "pie"])
def test_chart_types(kind):
    args = one_layout("chart")
    args["slides"][0]["chart"]["type"] = kind
    _, _, result = build(args)
    assert audit_pptx(io.BytesIO(result.data)).charts == 1


def test_slide_number_field_and_fonts():
    _, _, result = build(ALL_LAYOUTS)
    xml = reopen(result).slides[3].shapes._spTree.xml
    assert 'type="slidenum"' in xml
    fonts, sizes = set(), []
    for slide in reopen(result).slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                for run in (r for p in shape.text_frame.paragraphs for r in p.runs):
                    fonts.add(run.font.name)
                    sizes.append((shape.name, run.font.size.pt))
    assert fonts <= {FONT}
    assert all(size >= BODY_FLOOR for name, size in sizes if name not in ("Footer", "Slide number"))


def test_margins_and_left_aligned_titles():
    _, _, result = build(ALL_LAYOUTS)
    for slide in reopen(result).slides:
        title = slide.shapes.title
        assert title.left / 914400 >= 0.5
        assert title.text_frame.paragraphs[0].alignment == 1  # PP_ALIGN.LEFT


# --- переполнение: подбор кегля и разбиение ----------------------------------------------------------


def test_stress_deck_has_no_overflow_overlap_or_out_of_bounds():
    deck, _, result = build(long_text_deck())
    report = audit_pptx(io.BytesIO(result.data))
    assert report.issues == [], [(i.slide, i.kind, i.detail) for i in report.issues]
    assert result.slide_count > 8  # 8 исходных слайдов (с титульным) разбились
    assert deck.stats["slides_split"] >= 4


def test_long_title_is_cut_and_full_text_goes_to_notes(caplog):
    caplog.set_level(logging.WARNING)
    args = one_layout("bullets")
    args["slides"][0]["title"] = "Очень длинный заголовок слайда " * 6
    deck, _, _ = build(args)
    assert deck.stats["title_truncated"] == 1
    s = deck.slides[-1]
    assert len(s.title) <= 80 and s.title.endswith("…")
    assert "Полный заголовок" in s.notes
    assert any("title_truncated" in r.message for r in caplog.records)  # лог warning


def test_bullets_over_five_are_split_with_part_markers():
    args = one_layout("bullets")
    args["slides"][0]["bullets"] = [f"Пункт номер {i}" for i in range(8)]
    deck, slides, _ = build(args)
    parts = [s for s in slides if s.layout == "bullets"]
    assert len(parts) == 2 and all(len(p.bullets) <= 5 for p in parts)
    assert parts[0].title.endswith("(1/2)") and parts[1].title.endswith("(2/2)")
    assert deck.stats["bullets_split"] == 1 and deck.stats["slides_split"] == 1


def test_big_table_splits_by_rows_and_repeats_header():
    args = one_layout("table")
    args["slides"][0]["table"]["rows"] = [[str(i), "значение"] for i in range(20)]
    deck, slides, _ = build(args)
    tables = [s for s in slides if s.layout == "table"]
    assert len(tables) >= 3 and all(t.table.headers == ["А", "Б"] for t in tables)
    assert sum(len(t.table.rows) for t in tables) == 20
    assert deck.stats["table_split"] == 1


def test_fit_size_prefers_large_and_floors_at_14():
    short = L.fit_size(["Короткий пункт"], 10.0, 4.0)
    assert short == 22
    tight = L.fit_size(["слово " * 40], 4.0, 1.2)
    assert tight is None or tight >= 14
    sizes = [L.fit_size(["слово " * n], 6.0, 2.5) or 0 for n in (5, 20, 40)]
    assert sizes == sorted(sizes, reverse=True)  # больше текста — не крупнее кегль


def test_title_fit_two_lines_and_never_below_32():
    pt, text = L.title_fit("Короткий", L.CONTENT_W)
    assert pt == 40 and text == "Короткий"
    pt, text = L.title_fit("слово " * 40, L.CONTENT_W)
    assert pt == 32 and text.endswith("…")


# --- нормализация ответа модели ----------------------------------------------------------------------


def test_unknown_layout_is_inferred_and_counted(caplog):
    caplog.set_level(logging.WARNING)
    deck = normalize(
        {
            "title": "T",
            "slides": [
                {
                    "layout": "fancy",
                    "title": "A",
                    "chart": {"categories": ["x"], "series": [{"name": "s", "values": [1]}]},
                },
                {
                    "layout": "???",
                    "title": "B",
                    "stats": [{"value": "1", "label": "l"}, {"value": "2", "label": "m"}],
                },
                {"layout": "", "title": "C", "bullets": ["p"]},
            ],
        }
    )
    assert [s.layout for s in deck.slides[1:]] == ["chart", "stat", "bullets"]
    assert deck.stats["unknown_layout"] == 3
    assert sum("unknown_layout" in r.message for r in caplog.records) == 3


def test_layouts_without_data_fall_back():
    deck = normalize(
        {
            "title": "T",
            "slides": [
                {"layout": "chart", "title": "A", "bullets": ["b"]},
                {"layout": "table", "title": "B"},
                {"layout": "image", "title": "C", "bullets": ["b"]},  # без image_prompt
                {"layout": "timeline", "title": "D", "steps": [{"title": "один", "text": "шаг"}]},
            ],
        }
    )
    assert [s.layout for s in deck.slides[1:]] == ["bullets"] * 4
    assert deck.stats["layout_fallback"] == 4
    assert any("один: шаг" in b for b in deck.slides[4].bullets)


def test_garbage_input_never_raises():
    for args in (
        {},
        None,
        {"slides": "oops"},
        {"slides": [1, None, "x", {"layout": 5}]},
        {"title": 3, "theme": 9},
    ):
        deck = normalize(args)  # type: ignore[arg-type]
        assert deck.slides and deck.slides[0].layout == "title"
        deck2, slides = prepare_deck(args)  # type: ignore[arg-type]
        assert render(deck2, slides, {}).slide_count == len(slides)


def test_numbers_and_chart_lengths_are_coerced():
    deck = normalize(
        {
            "title": "T",
            "slides": [
                {
                    "layout": "chart",
                    "title": "A",
                    "chart": {
                        "type": "bar",
                        "categories": ["a", "b", "c"],
                        "series": [{"name": "s", "values": ["1 200,5", "12%", "abc", 4]}],
                    },
                }
            ],
        }
    )
    chart = deck.slides[1].chart
    assert chart.series[0][1] == [1200.5, 12.0, 0.0]  # лишнее обрезано, мусор -> 0
    assert deck.stats["bad_number"] == 1 and deck.stats["chart_values_length_fixed"] == 1


def test_pie_keeps_one_series_and_caps_apply():
    many = {
        "type": "pie",
        "categories": [str(i) for i in range(20)],
        "series": [{"name": f"s{i}", "values": list(range(20))} for i in range(6)],
    }
    deck = normalize({"title": "T", "slides": [{"layout": "chart", "title": "A", "chart": many}]})
    chart = deck.slides[1].chart
    assert len(chart.categories) == 12 and len(chart.series) == 1
    assert deck.stats["chart_categories_truncated"] == 1


def test_title_slide_is_never_duplicated():
    with_title = normalize(
        {
            "title": "T",
            "slides": [
                {"layout": "title", "title": "T"},
                {"layout": "bullets", "title": "B", "bullets": ["x"]},
            ],
        }
    )
    assert [s.layout for s in with_title.slides] == ["title", "bullets"]
    without = normalize(
        {
            "title": "T",
            "subtitle": "sub",
            "slides": [{"layout": "bullets", "title": "B", "bullets": ["x"]}],
        }
    )
    assert [s.layout for s in without.slides] == ["title", "bullets"] and without.slides[
        0
    ].subtitle == "sub"
    stray = normalize(
        {
            "title": "T",
            "slides": [
                {"layout": "bullets", "title": "B", "bullets": ["x"]},
                {"layout": "title", "title": "Again"},
            ],
        }
    )
    assert [s.layout for s in stray.slides] == ["title", "bullets"]  # < 12 слайдов: без разделов
    assert stray.stats["extra_title_slide"] == 1 and stray.stats["section_dropped"] == 1


def test_unknown_theme_defaults_to_graphite_and_theme_applies():
    assert normalize({"title": "T", "theme": "neon", "slides": []}).theme == "graphite"
    _, _, result = build({**one_layout("bullets"), "theme": "forest"})
    bg = reopen(result).slides[0].background.fill.fore_color.rgb
    assert str(bg) == PALETTES["forest"].dark.upper()


def test_slide_cap_enforced():
    args = {
        "title": "T",
        "slides": [
            {"layout": "section", "title": f"S{i}", "notes": "n"} for i in range(MAX_SLIDES + 40)
        ],
    }
    deck, slides = prepare_deck(args)
    assert len(slides) <= MAX_SLIDES and deck.stats["slides_over_cap"] == 1


def test_legacy_build_presentation_signature_still_works():
    data = build_presentation(
        "Название", "Подзаголовок", [{"title": "Слайд", "content": ["а", "б"]}]
    )
    prs = Presentation(io.BytesIO(data))
    assert len(prs.slides) == 2 and prs.slides[1].shapes.title.text == "Слайд"


# --- схема инструмента для трёх провайдеров ----------------------------------------------------------------


def _walk(node):
    yield node
    if isinstance(node, dict):
        for v in node.values():
            yield from _walk(v)
    elif isinstance(node, list):
        for v in node:
            yield from _walk(v)


def test_tool_schema_is_provider_neutral_and_compact():
    params = PRESENTATION_TOOL["parameters"]
    for node in _walk(params):
        if isinstance(node, dict):
            assert not set(node) & {
                "anyOf",
                "oneOf",
                "allOf",
                "$ref",
                "default",
                "additionalProperties",
                "nullable",
            }, node
            if "type" in node and not isinstance(
                node["type"], dict
            ):  # dict = свойство с именем type
                assert isinstance(node["type"], str)  # без union-типов вида ["string","null"]
    assert set(params["properties"]["slides"]["items"]["properties"]["layout"]["enum"]) == set(
        LAYOUTS
    )
    assert params["properties"]["theme"]["enum"] == list(PALETTES)
    # регрессия размера: схема платится в каждом запросе (~1,4k токенов ≈ 4,3k символов JSON)
    assert len(json.dumps(PRESENTATION_TOOL, ensure_ascii=False, separators=(",", ":"))) < 4800


def test_tool_schema_converts_for_gemini_openai_anthropic():
    decl = types.FunctionDeclaration(
        name=PRESENTATION_TOOL["name"],
        description=PRESENTATION_TOOL["description"],
        parameters=PRESENTATION_TOOL["parameters"],
    )
    assert "slides" in decl.parameters.properties  # Gemini (pydantic Schema) принимает схему
    anthropic_tool = {
        "name": PRESENTATION_TOOL["name"],
        "description": PRESENTATION_TOOL["description"],
        "input_schema": PRESENTATION_TOOL["parameters"],
    }
    openai_tool = {
        "type": "function",
        "strict": False,
        **{k: PRESENTATION_TOOL[k] for k in ("name", "description", "parameters")},
    }
    json.dumps(anthropic_tool), json.dumps(openai_tool)  # сериализуются без ошибок


# --- картинки -----------------------------------------------------------------------------------------------


CFG = {
    "max_images": 3,
    "hero_quality": "medium",
    "content_quality": "low",
    "hero_size": "1024x1536",
    "content_size": "1024x1024",
    "wide_size": "1536x1024",
    "concurrency": 2,
    "images_timeout_s": 0.5,
}


def deck_with_images(n_content: int):
    slides = [
        {
            "layout": "bullets",
            "title": f"S{i}",
            "bullets": ["x"],
            "image_prompt": f"prompt {i}",
            "notes": "n",
        }
        for i in range(n_content)
    ]
    deck, planned = prepare_deck(
        {"title": "T", "image_prompt": "hero prompt", "theme": "ocean", "slides": slides}
    )
    return deck, planned


def test_image_job_planning_respects_limits_and_priority(caplog):
    caplog.set_level(logging.WARNING)
    deck, planned = deck_with_images(6)
    jobs = deck_images.plan_image_jobs(deck, planned, CFG)
    assert [j[0] for j in jobs] == ["hero", 0, 1]  # hero первым, затем по порядку; max_images=3
    assert jobs[0][2:] == ("1024x1536", "medium") and jobs[1][2:] == ("1024x1024", "low")
    assert deck.stats["images_dropped"] == 4
    assert any("images_dropped" in r.message for r in caplog.records)
    deck, planned = deck_with_images(20)
    assert (
        len(deck_images.plan_image_jobs(deck, planned, {**CFG, "max_images": 99}))
        == DECK_MAX_IMAGES_HARD
    )  # жёсткий потолок 8


def test_image_prompt_only_where_a_layout_can_show_it():
    deck = normalize(
        {
            "title": "T",
            "slides": [
                {
                    "layout": "table",
                    "title": "A",
                    "table": {"headers": ["a"], "rows": [["1"]]},
                    "image_prompt": "x",
                }
            ],
        }
    )
    assert deck.stats["image_prompt_ignored"] == 1 and not deck.slides[1].image_prompt


async def test_failed_image_falls_back_to_layout_without_picture(monkeypatch):
    async def gen(prompt, size=None, quality=None):
        if "prompt 0" in prompt:
            raise RuntimeError("boom")
        return noisy_png(40, 40), "gpt-image-2.5-sunburst"

    monkeypatch.setattr(deck_images, "generate_image_bytes", gen)
    monkeypatch.setattr(deck_images, "load_config", lambda: {"deck": CFG})
    deck, planned = deck_with_images(2)
    images, model = await deck_images.generate_deck_images(deck, planned)
    assert set(images) == {"hero", 1} and model == "gpt-image-2.5-sunburst"
    assert deck.stats["image_failed"] == 1 and deck.stats["images_generated"] == 2
    result = render(deck, planned, images)  # слайд без картинки строится без исключений
    report = audit_pptx(io.BytesIO(result.data))
    assert report.issues == [] and report.pictures == 2  # hero + одна контентная


async def test_images_timeout_drops_slow_ones(monkeypatch):
    async def gen(prompt, size=None, quality=None):
        if "hero" in prompt:
            return noisy_png(40, 40), "m"
        await asyncio.sleep(5)
        return noisy_png(40, 40), "m"

    monkeypatch.setattr(deck_images, "generate_image_bytes", gen)
    monkeypatch.setattr(deck_images, "load_config", lambda: {"deck": CFG})
    deck, planned = deck_with_images(2)
    images, _ = await deck_images.generate_deck_images(deck, planned)
    assert set(images) == {"hero"} and deck.stats["image_timeout"] == 2


async def test_no_prompts_means_no_calls(monkeypatch):
    async def gen(*a, **k):
        raise AssertionError("не должно вызываться")

    monkeypatch.setattr(deck_images, "generate_image_bytes", gen)
    monkeypatch.setattr(deck_images, "load_config", lambda: {"deck": CFG})
    deck, planned = prepare_deck(one_layout("bullets"))
    assert await deck_images.generate_deck_images(deck, planned) == ({}, None)


def test_pictures_are_cropped_not_stretched_and_have_alt_text():
    src = noisy_png(300, 200)  # 3:2
    for box in ((0, 0, 4.9, 4.85), (0, 0, 5.133, 7.5), (0, 0, 6.3, 4.85)):
        out = Image.open(io.BytesIO(prepare_picture(src, box[2], box[3])))
        assert abs(out.width / out.height - box[2] / box[3]) < 0.02
    deck, planned = deck_with_images(1)
    result = render(deck, planned, {"hero": noisy_png(80, 120), 0: noisy_png(60, 60)})
    prs = reopen(result)
    pics = [sh for s in prs.slides for sh in s.shapes if sh.shape_type == 13]
    assert len(pics) == 2
    alts = {sh._element.nvPicPr.cNvPr.get("descr") for sh in pics}
    assert alts == {"hero prompt", "prompt 0"}  # alt-текст из image_prompt
    for sh in pics:
        assert len(sh.image.blob) < 200_000  # пережато в JPEG


def test_table_columns_never_split_words():
    table = L.Table(
        ["Модуль", "Назначение"],
        [["Аналитика", "Дашборды и отчёты"], ["Производственное", "Учёт выпуска"]],
    )
    for pt in (20, 16, 14):
        widths, _heights = L.table_geometry(table, pt)
        assert abs(sum(widths) - L.CONTENT_W) < 0.01
        for c, header in enumerate(table.headers):
            cells = [header] + [r[c] for r in table.rows]
            for text in cells:
                for word in text.split():
                    assert L.text_width(word, pt, bold=True) <= widths[c] - 0.24, (pt, word)
