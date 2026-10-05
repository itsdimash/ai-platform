import io
import logging
from datetime import date

import pdfplumber
import pytest
from docx import Document
from openpyxl import load_workbook

from app import limits
from app.utils import xlsx_builder
from app.utils.docx_builder import DOCUMENT_TOOL, build_document
from app.utils.pdf_builder import PDF_TOOL, build_pdf
from app.utils.xlsx_builder import SPREADSHEET_TOOL, build_spreadsheet, is_safe_formula, parse_cell

SECTIONS = [
    {
        "heading": "Цель",
        "paragraphs": ["Перенести склад на новый адрес до 1 декабря."],
        "bullets": ["Снизить аренду на 15%", "Ускорить отгрузку"],
    },
    {
        "heading": "План",
        "numbered": ["Подготовить помещение", "Перевезти стеллажи"],
        "table": {
            "headers": ["Этап", "Срок", "Бюджет"],
            "rows": [["Подготовка", "15.11", "1 200 000"], ["Перевозка", "25.11", "300 000"]],
        },
    },
    {"heading": "Повтор", "numbered": ["Первый", "Второй"]},
]


# --- docx ------------------------------------------------------------------------------------


def test_docx_structure_lists_table_and_page_numbers():
    doc = Document(io.BytesIO(build_document("Служебная записка", SECTIONS, "Для: директора")))
    styles = [p.style.name for p in doc.paragraphs]
    assert styles[0] == "Title" and doc.paragraphs[0].text == "Служебная записка"
    assert "Для: директора" in [p.text for p in doc.paragraphs]
    assert styles.count("Heading 1") == 3 and "List Bullet" in styles and "List Number" in styles
    assert len(doc.tables) == 1
    table = doc.tables[0]
    assert [c.text for c in table.rows[0].cells] == ["Этап", "Срок", "Бюджет"]
    assert "w:tblHeader" in table.rows[0]._tr.xml  # шапка повторяется на страницах
    assert 'w:fill="3A4654"' in table.rows[0].cells[0]._tc.xml  # заливка шапки
    footer_xml = doc.sections[0].footer.paragraphs[0]._p.xml
    assert "PAGE" in footer_xml and "NUMPAGES" in footer_xml  # «Стр. X из Y»
    assert doc.core_properties.title == "Служебная записка"


def test_docx_numbered_lists_restart_numbering():
    doc = Document(io.BytesIO(build_document("T", SECTIONS)))
    ids = [p._p.pPr.numPr.numId.val for p in doc.paragraphs if p.style.name == "List Number"]
    assert ids[0] == ids[1] and ids[2] == ids[3] and ids[0] != ids[2]  # второй список — с единицы


def test_docx_cyrillic_and_robust_to_odd_input():
    data = build_document(
        "Договор №1",
        [{"heading": "", "paragraphs": ["Текст"]}, {}, {"table": {"headers": [], "rows": []}}],
    )
    assert Document(io.BytesIO(data)).paragraphs[0].text == "Договор №1"


# --- pdf -------------------------------------------------------------------------------------


def test_pdf_page_numbers_lists_tables_and_cyrillic():
    many = [
        {
            **SECTIONS[0],
            "paragraphs": ["Длинный абзац про склад и логистику. " * 15],
            "heading": f"Раздел {i}",
        }
        for i in range(8)
    ]
    data = build_pdf("Отчёт о переносе склада", [*many, *SECTIONS], "Октябрь 2026")
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        pages = [p.extract_text() or "" for p in pdf.pages]
    total = len(pages)
    assert total >= 2
    for i, text in enumerate(pages, 1):
        assert f"Стр. {i} из {total}" in text
    joined = "\n".join(pages)
    assert "Отчёт о переносе склада" in joined and "Октябрь 2026" in joined
    assert "1. Подготовить помещение" in joined and "• Снизить аренду на 15%" in joined
    assert "Подготовка" in joined and "1 200 000" in joined


def test_pdf_escapes_markup():
    data = build_pdf("A & B <b>", [{"paragraphs": ["<script>x</script> & текст"]}])
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        text = pdf.pages[0].extract_text()
    assert "<script>x</script>" in text and "A & B <b>" in text


# --- xlsx ------------------------------------------------------------------------------------


@pytest.mark.parametrize(
    ("raw", "value", "fmt"),
    [
        ("1 234,50", 1234.5, "#,##0.00"),
        ("1234", 1234, "#,##0"),
        ("12", 12, "0"),
        ("12%", 0.12, "0%"),
        ("12,5 %", 0.125, "0.0%"),
        ("5 000 ₸", 5000, '#,##0 "₸"'),
        ("(1 200)", -1200, "#,##0"),
        ("31.12.2026", date(2026, 12, 31), "DD.MM.YYYY"),
        ("2026-01-05", date(2026, 1, 5), "DD.MM.YYYY"),
    ],
)
def test_parse_cell_numbers_dates_and_formats(raw, value, fmt):
    assert parse_cell(raw) == (value, fmt)


@pytest.mark.parametrize(
    "raw", ["Алматы", "007", "+77012345678", "1234567890123456789", "12abc", "", "32.13.2026"]
)
def test_identifiers_and_text_stay_text(raw):
    assert isinstance(parse_cell(raw)[0], str)


def test_formula_whitelist():
    for ok in (
        "=SUM(B2:B11)",
        "=AVERAGE(C2:C9)",
        "=ROUND(B2,2)",
        "=B2*C2",
        "=B2/C2+10",
        "=(B2+C2)*2",
    ):
        assert is_safe_formula(ok), ok
    for bad in (
        '=HYPERLINK("http://x")',
        "=cmd|' /C calc'!A0",
        "=EVIL(1)",
        "=SUM(B2:B11)+EVIL(1)",
        "=B2^2",
        '=IMPORTDATA("x")',
    ):
        assert not is_safe_formula(bad), bad


def make_sheet(**extra):
    return {
        "name": "Бюджет",
        "headers": ["Статья", "План, ₸", "Факт, ₸", "Исполнение"],
        "rows": [
            ["Закупки", "1 200 000", "1 150 000", "95,8%"],
            ["Логистика", "300 000", "340 000", "113%"],
            ["Аренда", "450000", "450000", "100%"],
        ],
        **extra,
    }


def test_xlsx_numbers_numeric_header_frozen_widths_and_formats():
    wb = load_workbook(io.BytesIO(build_spreadsheet([make_sheet()])))
    ws = wb["Бюджет"]
    assert ws.freeze_panes == "A2" and ws.auto_filter.ref == "A1:D4"
    assert ws["A1"].font.bold and ws["A1"].fill.fgColor.rgb.endswith("3A4654")
    assert (
        ws["B2"].value == 1200000
        and isinstance(ws["B2"].value, int)
        and ws["B2"].number_format == "#,##0"
    )
    assert ws["D2"].value == pytest.approx(0.958) and ws["D2"].number_format == "0.0%"
    assert ws["A2"].value == "Закупки"
    assert ws.column_dimensions["A"].width >= 10 and ws.column_dimensions["B"].width >= 10


def test_xlsx_totals_skip_percent_columns_and_use_sum_formulas():
    ws = load_workbook(io.BytesIO(build_spreadsheet([make_sheet(totals=True)]))).active
    assert ws["A5"].value == "Итого"
    assert (
        ws["B5"].value == "=SUM(B2:B4)" and ws["C5"].value == "=SUM(C2:C4)" and ws["B5"].font.bold
    )
    assert ws["D5"].value is None  # проценты не суммируем


def test_xlsx_unsafe_formulas_become_text_safe_ones_stay_formulas():
    sheet = {
        "name": "T",
        "headers": ["a", "b"],
        "rows": [["=EVIL(1)", "=SUM(B3:B3)"], ['=HYPERLINK("x")', "5"]],
    }
    ws = load_workbook(io.BytesIO(build_spreadsheet([sheet]))).active
    assert ws["A2"].data_type == "s" and ws["A3"].data_type == "s"
    assert ws["B2"].data_type == "f"


def test_xlsx_one_simple_chart():
    sheet = make_sheet(
        chart={"type": "bar", "title": "План и факт", "category_column": 1, "value_columns": [2, 3]}
    )
    ws = load_workbook(io.BytesIO(build_spreadsheet([sheet]))).active
    assert len(ws._charts) == 1
    assert len(load_workbook(io.BytesIO(build_spreadsheet([make_sheet()]))).active._charts) == 0
    bad = make_sheet(chart={"type": "pie", "category_column": 9, "value_columns": [99]})
    assert (
        len(load_workbook(io.BytesIO(build_spreadsheet([bad]))).active._charts) == 0
    )  # мусор игнорируется


def test_xlsx_row_cap_and_sheet_names(monkeypatch, caplog):
    caplog.set_level(logging.WARNING)
    monkeypatch.setattr(xlsx_builder, "MAX_XLSX_ROWS", 5)
    assert limits.MAX_XLSX_ROWS == 50_000
    sheet = {
        "name": "Плохое/имя:[лист]?" + "x" * 40,
        "headers": ["n"],
        "rows": [[str(i)] for i in range(50)],
    }
    wb = load_workbook(io.BytesIO(build_spreadsheet([sheet, {**sheet}])))
    assert wb.worksheets[0].max_row == 6  # шапка + 5 строк
    names = [w.title for w in wb.worksheets]
    assert (
        len(names[0]) <= 31
        and not set(names[0]) & set("[]:*?/\\")
        and len({n.lower() for n in names}) == 2
    )
    assert any("rows_truncated" in r.message for r in caplog.records)


def test_xlsx_never_leaves_workbook_without_sheets():
    assert load_workbook(io.BytesIO(build_spreadsheet([]))).worksheets


# --- схемы инструментов ----------------------------------------------------------------------------


@pytest.mark.parametrize("tool", [DOCUMENT_TOOL, PDF_TOOL, SPREADSHEET_TOOL])
def test_document_tools_have_summary_and_neutral_schemas(tool):
    props = tool["parameters"]["properties"]
    assert "summary" in props and "summary" in tool["parameters"]["required"]

    def walk(n):
        if isinstance(n, dict):
            assert not set(n) & {"anyOf", "oneOf", "$ref", "default", "additionalProperties"}
            for v in n.values():
                walk(v)
        elif isinstance(n, list):
            for v in n:
                walk(v)

    walk(tool["parameters"])


def test_xlsx_formula_cells_inherit_number_format_and_enter_totals():
    sheet = {
        "name": "Бюджет",
        "headers": ["Статья", "Окт", "Ноя", "Итого"],
        "rows": [["А", "1 000", "2 000", "=SUM(B2:C2)"], ["Б", "3 000", "4 000", "=SUM(B3:C3)"]],
        "totals": True,
    }
    ws = load_workbook(io.BytesIO(build_spreadsheet([sheet]))).active
    assert ws["D2"].number_format == "#,##0"  # формат формулы, а не General
    assert (
        ws["D4"].value == "=SUM(D2:D3)" and ws["D4"].number_format == "#,##0"
    )  # колонка формул — в итогах
