import io

from pptx import Presentation

from app.utils.tool_schema import SUMMARY_PROP

PRESENTATION_TOOL = {
    "name": "generate_presentation",
    "description": "Generates a PowerPoint presentation (.pptx) file and delivers it to the user as a downloadable attachment when requested.",
    "parameters": {
        "type": "object",
        "properties": {
            "summary": SUMMARY_PROP,
            "title": {"type": "string", "description": "The presentation main title"},
            "subtitle": {
                "type": "string",
                "description": "Presentation subtitle or author context",
            },
            "slides": {
                "type": "array",
                "description": "List of slide topics and content",
                "items": {
                    "type": "object",
                    "properties": {
                        "title": {"type": "string", "description": "Slide heading"},
                        "content": {
                            "type": "array",
                            "items": {"type": "string"},
                            "description": "Bullet points for the slide",
                        },
                    },
                    "required": ["title", "content"],
                },
            },
        },
        "required": ["title", "slides", "summary"],
    },
}


def build_presentation(title: str, subtitle: str, slides_data: list) -> bytes:
    """Собирает .pptx и возвращает байты (загрузка в R2 — в app/tools)."""
    prs = Presentation()

    title_slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(title_slide_layout)
    title_shape = slide.shapes.title
    subtitle_shape = slide.placeholders[1]
    title_shape.text = title
    subtitle_shape.text = subtitle or "Kerneu Group"

    bullet_slide_layout = prs.slide_layouts[1]
    for s_data in slides_data:
        slide = prs.slides.add_slide(bullet_slide_layout)
        shapes = slide.shapes
        title_shape = shapes.title
        body_shape = shapes.placeholders[1]

        title_shape.text = s_data.get("title", "")
        tf = body_shape.text_frame

        content_items = s_data.get("content", [])
        if content_items:
            tf.text = content_items[0]
            for item in content_items[1:]:
                p = tf.add_paragraph()
                p.text = item

    stream = io.BytesIO()
    prs.save(stream)
    return stream.getvalue()
