"""Извлечение ячейки тендерного .docx по её заголовку.

Нужная ячейка — с заголовком «Описание требуемых функциональных, технических, качественных,
эксплуатационных и иных характеристик закупаемого товара». Раскладки:
  A) заголовок в шапке колонки, содержимое — в ячейках ниже той же колонки;
  B) заголовок слева, содержимое — в той же строке справа.
Перебираются все таблицы документа, включая вложенные. Объединённые ячейки (gridSpan/vMerge)
python-docx отдаёт повторяющимися в row.cells — дедуплицируем по идентичности элемента _tc.
Заголовок ищется с нормализацией (регистр, пробелы, неразрывные пробелы, переносы) и по
ключевым словам, поэтому небольшие отличия формулировки не мешают.

Тихого отката на «весь документ» нет: если ячейка не найдена — CellNotFoundError.
"""

import io
import re
from dataclasses import dataclass, field

import docx
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph

# Ключевые слова заголовка (после нормализации). Нужны все группы.
HEADER_KEYWORDS = (("описание", "требуем"), ("характеристик",))
MAX_HEADER_CHARS = 400  # длиннее — это содержимое, а не заголовок
SHORT_HEADER_CHARS = 120  # «короткая» соседняя ячейка шапки

_W_P = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}p"
_W_TBL = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tbl"
_SPACES = re.compile(r"\s+")
_PUNCT = re.compile(r"[\"'«»“”„:;]")


class CellNotFoundError(Exception):
    """В документе нет ячейки с нужным заголовком (или она пуста)."""


@dataclass
class CellExtraction:
    text: str
    warnings: list[str] = field(default_factory=list)
    headers_matched: int = 0
    cells_used: int = 0


def normalize(text: str) -> str:
    text = (text or "").replace("\xa0", " ").replace("\u200b", "")
    text = _PUNCT.sub(" ", text.casefold())
    return _SPACES.sub(" ", text).strip()


def is_target_header(text: str) -> bool:
    norm = normalize(text)
    if not norm or len(norm) > MAX_HEADER_CHARS:
        return False
    return all(any(word in norm for word in group) for group in HEADER_KEYWORDS)


def cell_text(cell: _Cell) -> str:
    """Текст ячейки в порядке документа: абзацы и вложенные таблицы (строки через « | »)."""
    parts: list[str] = []
    for child in cell._tc.iterchildren():
        if child.tag == _W_P:
            parts.append(Paragraph(child, cell).text)
        elif child.tag == _W_TBL:
            for row in _rows(Table(child, cell)):
                cells = [cell_text(c).strip() for c in _unique(row)]
                if any(cells):
                    parts.append(" | ".join(cells))
    return "\n".join(p for p in parts if p.strip()).strip()


def _unique(cells: list[_Cell]) -> list[_Cell]:
    """Ячейки строки без повторов объединённых (по идентичности _tc), порядок сохранён."""
    seen: list = []
    result: list[_Cell] = []
    for cell in cells:
        if not any(cell._tc is tc for tc in seen):
            seen.append(cell._tc)
            result.append(cell)
    return result


def _rows(table: Table) -> list[list[_Cell]]:
    """Строки таблицы как списки ячеек по колонкам сетки (объединённые повторяются)."""
    rows: list[list[_Cell]] = []
    for index, row in enumerate(table.rows):
        try:
            rows.append(list(row.cells))
        except (ValueError, IndexError, KeyError):
            # нерегулярная сетка (gridBefore/gridAfter и т.п.) — берём ячейки как есть
            rows.append([_Cell(tc, table) for tc in row._tr.tc_lst])
        del index
    return rows


def _iter_tables(tables: list[Table], prefix: str = ""):
    """(таблица, путь вида «таблица 2 → ячейка (3,1) → вложенная таблица 1»)."""
    for number, table in enumerate(tables, 1):
        label = f"вложенная таблица {number}" if prefix else f"таблица {number}"
        path = f"{prefix} → {label}" if prefix else label
        yield table, path
        for r, row in enumerate(_rows(table), 1):
            for c, cell in enumerate(row, 1):
                if c > 1 and cell._tc is row[c - 2]._tc:
                    continue  # повтор объединённой ячейки
                yield from _iter_tables(cell.tables, f"{path} → ячейка ({r},{c})")


@dataclass
class CellMatch:
    path: str  # где найдено: «таблица 1» или «таблица 2 → ячейка (1,1) → вложенная таблица 1»
    row: int  # строка заголовка (с 1)
    col: int  # колонка сетки заголовка (с 1)
    layout: str  # "A" — заголовок в шапке колонки, "B" — заголовок слева
    text: str


def _find_in_table(table: Table, path: str) -> list[CellMatch]:
    """Найденные в таблице заголовки с содержимым."""
    grid = _rows(table)
    found: list[CellMatch] = []
    claimed: list = []  # _tc уже учтённых заголовков
    for r, row in enumerate(grid):
        for c, cell in enumerate(row):
            if not any(cell._tc is tc for tc in claimed) and is_target_header(cell_text(cell)):
                claimed.append(cell._tc)
                text, layout = _content_for(grid, r, c, cell)
                if text:
                    found.append(CellMatch(path, r + 1, c + 1, layout, text))
    return found


def _content_for(grid: list[list[_Cell]], r: int, c: int, header: _Cell) -> tuple[str, str]:
    """(текст, раскладка). Раскладка выбирается по длине текста: содержимое описания —
    длинное, а подписи строк и соседние заголовки шапки — короткие."""
    row = grid[r]
    right = [x for x in _unique(row) if x._tc is not header._tc and _is_right_of(row, header, x)]
    right_texts = [t for t in (cell_text(x) for x in right) if t.strip()]
    below = _below(grid, r, c, header)

    right_long = any(len(t) > SHORT_HEADER_CHARS for t in right_texts)
    below_long = any(len(t) > SHORT_HEADER_CHARS for t in below)
    if right_long:
        return "\n".join(right_texts), "B"  # заголовок слева, содержимое справа
    if below_long or (below and not right_texts):
        return "\n".join(below), "A"  # шапка колонки
    if right_texts and len(_unique(row)) == 2:  # короткая пара «подпись — значение»
        return "\n".join(right_texts), "B"
    return "\n".join(below), "A"


def _is_right_of(row: list[_Cell], header: _Cell, other: _Cell) -> bool:
    header_idx = next(i for i, x in enumerate(row) if x._tc is header._tc)
    other_idx = next(i for i, x in enumerate(row) if x._tc is other._tc)
    return other_idx > header_idx


def _below(grid: list[list[_Cell]], r: int, c: int, header: _Cell) -> list[str]:
    texts: list[str] = []
    seen: list = [header._tc]
    for row in grid[r + 1 :]:
        if c >= len(row):
            continue
        cell = row[c]
        if any(cell._tc is tc for tc in seen):
            continue  # продолжение вертикального объединения / повтор
        seen.append(cell._tc)
        text = cell_text(cell)
        if text.strip() and not is_target_header(text):
            texts.append(text)
    return texts


def find_description_cells(data: bytes) -> list[CellMatch]:
    """Все найденные ячейки с описанием характеристик (в порядке документа)."""
    document = docx.Document(io.BytesIO(data))
    matches: list[CellMatch] = []
    for table, path in _iter_tables(list(document.tables)):
        matches.extend(_find_in_table(table, path))
    return matches


def extract_description_cell(data: bytes) -> CellExtraction:
    """Текст ячейки «Описание требуемых … характеристик …» из .docx (байты).

    Несколько подходящих ячеек — тексты склеиваются по порядку документа, добавляется warning.
    Нет ни одной (или содержимое пусто) — CellNotFoundError."""
    texts = [m.text for m in find_description_cells(data)]
    headers = len(texts)
    if not texts:
        raise CellNotFoundError(
            "В документе не найдена ячейка таблицы с заголовком «Описание требуемых "
            "функциональных, технических, качественных, эксплуатационных и иных характеристик "
            "закупаемого товара»."
        )
    warnings: list[str] = []
    if len(texts) > 1:
        warnings.append(
            f"Найдено несколько ячеек с описанием характеристик ({len(texts)}): "
            "их тексты объединены по порядку документа."
        )
    return CellExtraction(
        text="\n\n".join(texts), warnings=warnings, headers_matched=headers, cells_used=len(texts)
    )
