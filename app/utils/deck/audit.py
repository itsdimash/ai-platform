"""Программная проверка готового .pptx: границы слайда, наложения текстовых блоков,
переполнение текста (по той же оценке ширины, что использует планировщик), минимальные
кегли, наличие заметок. Используется тестами и верификацией; на пользователя не влияет."""

import math
from dataclasses import dataclass, field

from pptx import Presentation

from app.utils.deck import layout as L
from app.utils.deck.schema import has_placeholder
from app.utils.deck.themes import BODY_FLOOR, CAPTION

EMU = 914400
TOL = 0.02  # дюймов
SPARSE_MIN = 0.55  # доля области контента, которую должен занимать блок
_CHROME = ("Footer", "Slide number", "Footnote", "Title accent")


@dataclass
class Issue:
    slide: int
    kind: str  # bounds | overlap | overflow | small_font | word_overflow | placeholder | sparse
    detail: str


@dataclass
class AuditReport:
    slides: int
    issues: list[Issue] = field(default_factory=list)
    warnings: list[Issue] = field(default_factory=list)  # не дефекты: sparse (контента мало)
    notes: int = 0
    tables: int = 0
    charts: int = 0
    pictures: int = 0

    def of(self, kind: str) -> list[Issue]:
        return [i for i in self.issues if i.kind == kind]


def _box(shape) -> tuple[float, float, float, float]:
    return shape.left / EMU, shape.top / EMU, shape.width / EMU, shape.height / EMU


def _paragraphs(shape):
    """[(runs[(text, pt, bold)], bullet, space_after_in)] непустых абзацев."""
    for para in shape.text_frame.paragraphs:
        runs = [
            (r.text, r.font.size.pt if r.font.size else 18, bool(r.font.bold))
            for r in para.runs
            if r.text
        ]
        if not runs:
            continue
        bullet = (
            para._p.pPr is not None
            and para._p.pPr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}buChar")
            is not None
        )
        after = para.space_after.pt / 72 if para.space_after is not None else None
        yield runs, bullet, after


def _paragraph_lines(runs, width: float) -> tuple[int, float]:
    """(число строк, наибольший кегль). Абзац из run'ов разного кегля (число + мелкая единица)
    считается по сумме ширин: он однострочный, пока суммарная ширина помещается."""
    biggest = max(size for _, size, _ in runs)
    if len({size for _, size, _ in runs}) == 1:
        text = "".join(t for t, _, _ in runs)
        return len(L.wrap_lines(text, biggest, width, any(b for _, _, b in runs))), biggest
    total = sum(L.text_width(t, size, bold) for t, size, bold in runs)
    return max(1, math.ceil(total / width)), biggest


def _needed_height(shape) -> float:
    width = shape.width / EMU
    total = 0.0
    paras = list(_paragraphs(shape))
    for count, (runs, bullet, after) in enumerate(paras):
        lines, size = _paragraph_lines(runs, width - (L.BULLET_INDENT if bullet else 0))
        total += lines * L.line_height(size)
        if count < len(paras) - 1:
            total += after if after is not None else L.para_gap(size)
    return total


def _word_overflow(shape) -> list[str]:
    """Самое длинное слово каждого run'а должно помещаться в ширину рамки (минус маркер)."""
    width = shape.width / EMU
    out = []
    for runs, bullet, _ in _paragraphs(shape):
        avail = width - (L.BULLET_INDENT if bullet else 0)
        for text, size, bold in runs:
            for word in L.words(text):
                w = L.text_width(word, size, bold)
                if w > avail + TOL:
                    out.append(f"{word!r} {size:g} pt: {w:.2f} > {avail:.2f} in")
    return out


def _table_word_overflow(shape) -> list[str]:
    out = []
    tbl = shape.table
    for r, row in enumerate(tbl.rows):
        for c, cell in enumerate(row.cells):
            width = tbl.columns[c].width / EMU - (cell.margin_left + cell.margin_right) / EMU
            for para in cell.text_frame.paragraphs:
                for run in para.runs:
                    size = run.font.size.pt if run.font.size else 18
                    for word in L.words(run.text):
                        w = L.text_width(word, size, bool(run.font.bold))
                        if w > width + TOL:
                            out.append(f"r{r}c{c} {word!r}: {w:.2f} > {width:.2f} in")
    return out


def _intersection(a, b) -> float:
    x = max(0.0, min(a[0] + a[2], b[0] + b[2]) - max(a[0], b[0]))
    y = max(0.0, min(a[1] + a[3], b[1] + b[3]) - max(a[1], b[1]))
    return x * y


def _fill_ratio(slide) -> float | None:
    """Доля области контента, занятая блоком (None — для тёмных слайдов без области контента)."""
    title = slide.shapes.title
    if title is None or abs(title.top / EMU - L.TITLE_BOX[1]) > 0.01:
        return None
    tops, bottoms = [], []
    for shape in slide.shapes:
        if shape == title or shape.name in _CHROME:
            continue
        _, y, _, h = _box(shape)
        if y >= L.CONTENT_Y - 0.05:
            tops.append(y)
            bottoms.append(min(y + h, L.CONTENT_BOTTOM))
    if not tops:
        return 0.0
    return (max(bottoms) - min(tops)) / L.CONTENT_H


def audit_pptx(data) -> AuditReport:
    prs = Presentation(data)
    sw, sh = prs.slide_width / EMU, prs.slide_height / EMU
    report = AuditReport(slides=len(prs.slides))
    for n, slide in enumerate(prs.slides, 1):
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            report.notes += 1
        text_shapes = []
        for shape in slide.shapes:
            x, y, w, h = _box(shape)
            if x < -TOL or y < -TOL or x + w > sw + TOL or y + h > sh + TOL:
                report.issues.append(
                    Issue(n, "bounds", f"{shape.name}: {x:.2f},{y:.2f} {w:.2f}x{h:.2f}")
                )
            if shape.shape_type is not None and shape.shape_type == 13:
                report.pictures += 1
            if shape.has_chart if hasattr(shape, "has_chart") else False:
                report.charts += 1
            if shape.has_table if hasattr(shape, "has_table") else False:
                report.tables += 1
                for detail in _table_word_overflow(shape):
                    report.issues.append(Issue(n, "word_overflow", f"{shape.name}: {detail}"))
                continue
            if shape.has_text_frame and shape.text_frame.text.strip():
                text_shapes.append(shape)
                for detail in _word_overflow(shape):
                    report.issues.append(Issue(n, "word_overflow", f"{shape.name}: {detail}"))
                if has_placeholder(shape.text_frame.text):
                    report.issues.append(
                        Issue(n, "placeholder", f"{shape.name}: {shape.text_frame.text[:40]!r}")
                    )
                need = _needed_height(shape)
                if need > h * 1.04 + 0.03:
                    report.issues.append(
                        Issue(
                            n,
                            "overflow",
                            f"{shape.name}: нужно {need:.2f} in, есть {h:.2f} in — {shape.text_frame.text[:40]!r}",
                        )
                    )
                for run in (r for p in shape.text_frame.paragraphs for r in p.runs):
                    if (
                        run.font.size
                        and run.font.size.pt < CAPTION
                        or run.font.size
                        and run.font.size.pt < BODY_FLOOR
                        and shape.name not in ("Footer", "Slide number", "Footnote")
                    ):
                        report.issues.append(
                            Issue(n, "small_font", f"{shape.name}: {run.font.size.pt} pt")
                        )
        if slide.has_notes_slide and has_placeholder(slide.notes_slide.notes_text_frame.text):
            report.issues.append(Issue(n, "placeholder", "notes"))
        sparse = _fill_ratio(slide)
        if sparse is not None and sparse < SPARSE_MIN:
            report.warnings.append(Issue(n, "sparse", f"контент занимает {sparse:.0%} области"))
        for i, a in enumerate(text_shapes):
            for b in text_shapes[i + 1 :]:
                ba, bb = _box(a), _box(b)
                inter = _intersection(ba, bb)
                if inter > 0.02 * min(ba[2] * ba[3], bb[2] * bb[3]):
                    report.issues.append(Issue(n, "overlap", f"{a.name} × {b.name}"))
    return report


__all__ = ["AuditReport", "Issue", "audit_pptx"]
