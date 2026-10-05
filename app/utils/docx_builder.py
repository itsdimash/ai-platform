"""Word-документы (.docx) по вызову инструмента generate_document.

Титульный блок, заголовки, абзацы, маркированные и нумерованные списки (нумерация каждого
списка начинается с 1), таблицы со стилем (шапка, зебра, повтор шапки на страницах) и
номера страниц «Стр. X из Y». Цвета — палитра graphite из deck.themes."""

import io
import re

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from app.utils.deck.themes import FONT, get_palette
from app.utils.tool_schema import SECTION_SCHEMA, SUMMARY_PROP

DOCUMENT_TOOL = {
    "name": "generate_document",
    "description": (
        "Create a Word .docx document: title block, sections with headings, paragraphs, bullet and "
        "numbered lists, tables. Use for reports, memos, letters, proposals, contract drafts. "
        "Write complete finished text in the user's language, not an outline."
    ),
    "parameters": {
        "type": "object",
        "properties": {
            "title": {"type": "string"},
            "subtitle": {"type": "string", "description": "Optional: date, author, addressee"},
            "summary": SUMMARY_PROP,
            "sections": {"type": "array", "items": SECTION_SCHEMA},
        },
        "required": ["title", "sections", "summary"],
    },
}

_NUMERIC = re.compile(r"^[\s +\-–−]?[\d\s .,]+\s?[%₸$€₽]?$")


def _rgb(hex6: str) -> RGBColor:
    return RGBColor.from_string(hex6.upper())


def _shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def _set_font(style, size: float, *, bold: bool | None = None, color: str | None = None) -> None:
    style.font.name = FONT
    style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    if color:
        style.font.color.rgb = _rgb(color)
    r_pr = style.element.get_or_add_rPr()
    fonts = r_pr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        r_pr.append(fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        fonts.set(qn(attr), FONT)


def _bottom_rule(paragraph, color: str) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "12")
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), color)
    borders.append(bottom)
    p_pr.append(borders)


def _field(paragraph, instruction: str) -> None:
    run = paragraph.add_run()
    for kind, text in (
        ("begin", None),
        (None, instruction),
        ("separate", None),
        (None, "1"),
        ("end", None),
    ):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        elif text == instruction:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = f" {instruction} "
        else:
            el = OxmlElement("w:t")
            el.text = text
        run._r.append(el)


def _footer_page_numbers(doc, muted: str) -> None:
    footer = doc.sections[0].footer
    para = footer.paragraphs[0]
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for chunk, field in (("Стр. ", None), (None, "PAGE"), (" из ", None), (None, "NUMPAGES")):
        if field:
            _field(para, field)
        else:
            para.add_run(chunk)
    for run in para.runs:
        run.font.size = Pt(9)
        run.font.name = FONT
        run.font.color.rgb = _rgb(muted)


def _new_numbering_instance(doc) -> int | None:
    """Создаёт новый экземпляр нумерации для стиля List Number, чтобы нумерация очередного
    списка начиналась с 1 (иначе все списки документа нумеруются сквозно)."""
    style = doc.styles["List Number"]
    num_pr = style.element.pPr.numPr if style.element.pPr is not None else None
    if num_pr is None or num_pr.numId is None:
        return None
    numbering = doc.part.numbering_part.numbering_definitions._numbering
    base_id = num_pr.numId.val
    abstract = None
    for num in numbering.findall(qn("w:num")):
        if num.get(qn("w:numId")) == str(base_id):
            abstract = num.find(qn("w:abstractNumId")).get(qn("w:val"))
    if abstract is None:
        return None
    new_id = max(int(n.get(qn("w:numId"))) for n in numbering.findall(qn("w:num"))) + 1
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(new_id))
    ref = OxmlElement("w:abstractNumId")
    ref.set(qn("w:val"), abstract)
    num.append(ref)
    override = OxmlElement("w:lvlOverride")
    override.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:startOverride")
    start.set(qn("w:val"), "1")
    override.append(start)
    num.append(override)
    numbering.append(num)
    return new_id


def _add_table(doc, headers: list[str], rows: list[list[str]], palette) -> None:
    ncols = max([len(headers)] + [len(r) for r in rows])
    if ncols == 0:
        return
    table = doc.add_table(rows=0, cols=ncols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    all_rows = ([headers] if headers else []) + rows
    weights = [
        max(6, min(40, max((len(r[c]) for r in all_rows if c < len(r)), default=6)))
        for c in range(ncols)
    ]
    usable = 16.6  # см: A4 минус поля
    for i, row in enumerate(all_rows):
        header = bool(headers) and i == 0
        cells = table.add_row().cells
        tr_pr = table.rows[-1]._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)
        if header:
            tr_pr.append(OxmlElement("w:tblHeader"))  # повтор шапки на каждой странице
        for c in range(ncols):
            text = row[c] if c < len(row) else ""
            cell = cells[c]
            cell.width = Cm(usable * weights[c] / sum(weights))
            para = cell.paragraphs[0]
            para.paragraph_format.space_after = Pt(2)
            para.paragraph_format.space_before = Pt(2)
            run = para.add_run(text)
            run.font.size = Pt(10.5)
            run.font.name = FONT
            if header:
                run.font.bold = True
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                _shade(cell, palette.primary)
            else:
                run.font.color.rgb = _rgb(palette.text)
                if (i - (1 if headers else 0)) % 2 == 0:
                    _shade(cell, palette.surface)
                if _NUMERIC.match(text or "x"):
                    para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def build_document(title: str, sections: list, subtitle: str = "") -> bytes:
    """Собирает .docx и возвращает байты (загрузка в R2 — в app/tools)."""
    p = get_palette("graphite")
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.2)
    sec.top_margin, sec.bottom_margin = Cm(2.0), Cm(2.0)

    _set_font(doc.styles["Normal"], 11)
    normal = doc.styles["Normal"].paragraph_format
    normal.space_after, normal.line_spacing = Pt(6), 1.15
    _set_font(doc.styles["Title"], 26, bold=True, color=p.primary)
    _set_font(doc.styles["Heading 1"], 16, bold=True, color=p.primary)
    _set_font(doc.styles["Heading 2"], 13, bold=True, color=p.primary)
    doc.styles["Heading 1"].paragraph_format.space_before = Pt(16)
    doc.styles["Heading 1"].paragraph_format.space_after = Pt(6)
    for name in ("List Bullet", "List Number"):
        _set_font(doc.styles[name], 11)
        doc.styles[name].paragraph_format.space_after = Pt(3)
    doc.core_properties.title = title
    doc.core_properties.author = "Kerneu AI"

    head = doc.add_paragraph(title, style="Title")
    if not subtitle:
        _bottom_rule(head, p.accent1)
    if subtitle:
        sub = doc.add_paragraph()
        run = sub.add_run(subtitle)
        run.font.size = Pt(12)
        run.font.color.rgb = _rgb(p.muted)
        _bottom_rule(sub, p.accent1)

    for section in sections:
        heading = (section.get("heading") or "").strip()
        if heading:
            doc.add_heading(heading, level=1)
        for text in section.get("paragraphs") or []:
            doc.add_paragraph(str(text))
        for text in section.get("bullets") or []:
            doc.add_paragraph(str(text), style="List Bullet")
        numbered = section.get("numbered") or []
        if numbered:
            num_id = _new_numbering_instance(doc)
            for text in numbered:
                para = doc.add_paragraph(str(text), style="List Number")
                if num_id is not None:
                    num_pr = para._p.get_or_add_pPr().get_or_add_numPr()
                    num_pr.get_or_add_numId().val = num_id
                    num_pr.get_or_add_ilvl().val = 0
        table = section.get("table")
        if isinstance(table, dict):
            headers = [str(h) for h in (table.get("headers") or [])]
            rows = [[str(c) for c in r] for r in (table.get("rows") or []) if isinstance(r, list)]
            _add_table(doc, headers, rows, p)

    _footer_page_numbers(doc, p.muted)
    stream = io.BytesIO()
    doc.save(stream)
    return stream.getvalue()
