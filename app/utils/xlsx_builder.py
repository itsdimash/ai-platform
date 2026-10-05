"""Excel-таблицы (.xlsx) по вызову инструмента generate_spreadsheet.

- ячейки приходят строками (так проще схема для всех провайдеров) и приводятся к типам:
  числа остаются числами (парсер понимает «1 234,50», «12%», «5 000 ₸», даты ДД.ММ.ГГГГ),
  идентификаторы с ведущими нулями и телефоны остаются текстом;
- жирная закреплённая шапка, автофильтр, ширина колонок, форматы чисел по колонкам;
- формулы допускаются только по белому списку (SUM/AVERAGE/MIN/MAX/COUNT/ROUND и простая
  арифметика над ячейками); любая другая строка с «=» записывается как текст;
- необязательные итоги («Итого» с SUM) и одна простая диаграмма на лист;
- потолок строк на лист — MAX_XLSX_ROWS (limits.py).
"""

import io
import logging
import re
from datetime import date

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.limits import MAX_XLSX_ROWS
from app.utils.deck.themes import FONT, get_palette
from app.utils.tool_schema import SUMMARY_PROP

logger = logging.getLogger(__name__)

SPREADSHEET_TOOL = {
    "name": "generate_spreadsheet",
    "description": (
        "Create an Excel .xlsx file. Cells are strings: write numbers plainly (1234.5, 12%, 5000), "
        "dates as DD.MM.YYYY. Realistic complete data (10+ rows unless told otherwise). "
        "Formulas only like =SUM(B2:B11). User's language."
    ),
    "parameters": {
        "type": "object",
        "properties": {
            "filename": {"type": "string", "description": "Base name without extension"},
            "summary": SUMMARY_PROP,
            "sheets": {
                "type": "array",
                "items": {
                    "type": "object",
                    "properties": {
                        "name": {"type": "string"},
                        "headers": {"type": "array", "items": {"type": "string"}},
                        "rows": {
                            "type": "array",
                            "items": {"type": "array", "items": {"type": "string"}},
                        },
                        "totals": {
                            "type": "boolean",
                            "description": "Add a totals row (SUM of numeric columns)",
                        },
                        "chart": {
                            "type": "object",
                            "description": "Optional one chart over this sheet's data",
                            "properties": {
                                "type": {"type": "string", "enum": ["bar", "line", "pie"]},
                                "title": {"type": "string"},
                                "category_column": {"type": "integer", "description": "1-based"},
                                "value_columns": {"type": "array", "items": {"type": "integer"}},
                            },
                        },
                    },
                    "required": ["name", "headers", "rows"],
                },
            },
        },
        "required": ["filename", "sheets", "summary"],
    },
}

# --- Разбор ячеек ---------------------------------------------------------------------------

_CURRENCY = r"(?:₸|тг|тенге|\$|€|₽|руб\.?|KZT|USD|EUR|RUB)"
_NUM_CORE = r"[+\-−–]?\d{1,3}(?:[  ]\d{3})*(?:[.,]\d+)?|[+\-−–]?\d+(?:[.,]\d+)?"
_NUMBER = re.compile(
    rf"^(?P<neg>[(])?\s*(?P<cur1>{_CURRENCY})?\s*(?P<num>{_NUM_CORE})\s*(?P<pct>%)?\s*(?P<cur2>{_CURRENCY})?\s*[)]?$",
    re.IGNORECASE,
)
_DATE_RU = re.compile(r"^(\d{1,2})\.(\d{1,2})\.(\d{4})$")
_DATE_ISO = re.compile(r"^(\d{4})-(\d{2})-(\d{2})$")
_CELL = r"[A-Z]{1,3}\d{1,6}"
_RANGE = rf"{_CELL}(?::{_CELL})?"
_FORMULA_FUNC = re.compile(
    rf"^=(?:SUM|AVERAGE|MIN|MAX|COUNT|ROUND)\({_RANGE}(?:,\s*\d{{1,2}})?\)$", re.IGNORECASE
)
_FORMULA_ARITH = re.compile(
    rf"^=\(?{_CELL}\)?(?:\s*[+\-*/]\s*(?:\(?{_CELL}\)?|\d+(?:\.\d+)?))*$", re.IGNORECASE
)


def is_safe_formula(text: str) -> bool:
    """Белый список формул: функции агрегации по диапазону или простая арифметика над ячейками."""
    return bool(_FORMULA_FUNC.match(text) or _FORMULA_ARITH.match(text))


def parse_cell(raw: object) -> tuple[object, str | None]:
    """(значение, number_format|None). Возвращает int/float/date/str/«=формула»."""
    text = "" if raw is None else str(raw).strip()
    if not text:
        return "", None
    if text.startswith("="):
        return (text, None) if is_safe_formula(text) else ("'" + text, "@")  # не формула -> текст
    iso, ru = _DATE_ISO.match(text), _DATE_RU.match(text)
    try:
        if iso:
            return date(int(iso[1]), int(iso[2]), int(iso[3])), "DD.MM.YYYY"
        if ru:
            return date(int(ru[3]), int(ru[2]), int(ru[1])), "DD.MM.YYYY"
    except ValueError:
        return text, None
    m = _NUMBER.match(text)
    if not m:
        return text, None
    digits = re.sub(r"[  +\-−–]", "", m["num"])
    # ведущие нули («007», артикулы), телефоны и слишком длинные числа — это идентификаторы
    if (len(digits.split(",")[0].split(".")[0]) > 15) or (
        re.match(r"^0\d", digits) and not re.match(r"^0[.,]", digits)
    ):
        return text, None
    if text.startswith("+") and len(digits) >= 10 and not m["pct"]:
        return text, None
    number = float(
        re.sub(r"[  ]", "", m["num"]).replace(",", ".").replace("−", "-").replace("–", "-")
    )
    if m["neg"]:
        number = -abs(number)
    cur = (m["cur1"] or m["cur2"] or "").strip()
    decimals = len(re.split(r"[.,]", m["num"])[1]) if re.search(r"[.,]\d+$", m["num"]) else 0
    if m["pct"]:
        number /= 100
        return number, "0.0%" if decimals else "0%"
    if cur:
        symbol = {"тг": "₸", "тенге": "₸", "kzt": "₸", "usd": "$", "eur": "€", "rub": "₽"}.get(
            cur.lower().rstrip("."), cur
        )
        symbol = "₽" if symbol.lower().startswith("руб") else symbol
        base = "#,##0.00" if decimals else "#,##0"
        return number, f'{base} "{symbol}"'
    value: object = int(number) if decimals == 0 and abs(number) < 1e15 else number
    fmt = (
        ("#,##0." + "0" * min(decimals, 4))
        if decimals
        else ("#,##0" if abs(number) >= 1000 else "0")
    )
    return value, fmt


def _column_format(values: list[tuple[object, str | None]]) -> str | None:
    fmts = [f for v, f in values if f and not isinstance(v, str)]
    if not fmts:
        return None
    return max(set(fmts), key=fmts.count)


def _safe_sheet_name(name: str, used: set[str]) -> str:
    cleaned = re.sub(r"[\[\]:*?/\\]", " ", name or "Лист").strip()[:31] or "Лист"
    base, n = cleaned, 2
    while cleaned.lower() in used:
        suffix = f" ({n})"
        cleaned, n = base[: 31 - len(suffix)] + suffix, n + 1
    used.add(cleaned.lower())
    return cleaned


def _chart(ws, spec: dict, header_row: int, first: int, last: int, ncols: int) -> None:
    kind = str(spec.get("type") or "bar").lower()
    cat_col = int(spec.get("category_column") or 1)
    value_cols = [
        int(c)
        for c in (spec.get("value_columns") or [])
        if 1 <= int(c) <= ncols and int(c) != cat_col
    ]
    if not value_cols or not 1 <= cat_col <= ncols or last < first:
        return
    chart = {"line": LineChart, "pie": PieChart}.get(kind, BarChart)()
    chart.title = str(spec.get("title") or "") or None
    chart.height, chart.width = 9, 18
    categories = Reference(ws, min_col=cat_col, min_row=first, max_row=last)
    for col in value_cols[: 1 if kind == "pie" else 4]:
        chart.add_data(
            Reference(ws, min_col=col, min_row=header_row, max_row=last), titles_from_data=True
        )
    chart.set_categories(categories)
    ws.add_chart(chart, f"{get_column_letter(ncols + 2)}{header_row}")


def build_spreadsheet(sheets: list) -> bytes:
    """Собирает .xlsx и возвращает байты (загрузка в R2 — в app/tools)."""
    palette = get_palette("graphite")
    header_fill = PatternFill("solid", fgColor=palette.primary)
    thin = Side(style="thin", color="C9D1D9")
    wb = Workbook()
    wb.remove(wb.active)  # убираем дефолтный пустой лист
    used: set[str] = set()

    for sheet_data in sheets:
        ws = wb.create_sheet(title=_safe_sheet_name(sheet_data.get("name"), used))
        headers = [str(h) for h in (sheet_data.get("headers") or [])]
        rows = [list(r) for r in (sheet_data.get("rows") or []) if isinstance(r, list | tuple)]
        if len(rows) > MAX_XLSX_ROWS:
            logger.warning("xlsx: rows_truncated %s -> %s", len(rows), MAX_XLSX_ROWS)
            rows = rows[:MAX_XLSX_ROWS]
        ncols = max([len(headers)] + [len(r) for r in rows] + [1])
        parsed = [[parse_cell(r[c] if c < len(r) else "") for c in range(ncols)] for r in rows]
        col_formats = [_column_format([row[c] for row in parsed]) for c in range(ncols)]
        # Формат «по умолчанию» для ячеек-формул: самый частый числовой формат листа
        # (колонка из одних формул, например итогов по строке, иначе осталась бы General).
        numeric_fmts = [f for f in col_formats if f and "%" not in f and "DD" not in f]
        sheet_fmt = max(set(numeric_fmts), key=numeric_fmts.count) if numeric_fmts else None
        col_formats = [
            f
            or (
                sheet_fmt
                if any(
                    isinstance(v, str) and v.startswith("=") for v, _ in (row[c] for row in parsed)
                )
                else None
            )
            for c, f in enumerate(col_formats)
        ]

        header_row = 1 if headers else 0
        if headers:
            ws.append((headers + [""] * ncols)[:ncols])
            for cell in ws[1]:
                cell.font = Font(name=FONT, bold=True, color="FFFFFF", size=11)
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            ws.row_dimensions[1].height = 24
            ws.freeze_panes = "A2"

        for row in parsed:
            ws.append([v for v, _ in row])
            r = ws.max_row
            for c, (value, fmt) in enumerate(row):
                cell = ws.cell(row=r, column=c + 1)
                cell.font = Font(name=FONT, size=11)
                cell.border = Border(bottom=thin)
                if isinstance(value, str) and value.startswith("'="):
                    cell.value = value[1:]
                    cell.data_type = "s"  # текст, не формула
                elif isinstance(value, str) and value.startswith("="):
                    cell.value = value
                if not isinstance(value, str) or value.startswith("="):
                    cell.number_format = fmt or col_formats[c] or "General"
                    cell.alignment = Alignment(horizontal="right", vertical="top")
                else:
                    cell.alignment = Alignment(vertical="top", wrap_text=len(value) > 60)

        first_data, last_data = header_row + 1, ws.max_row
        if sheet_data.get("totals") and rows:
            total_row = ws.max_row + 1
            ws.cell(row=total_row, column=1, value="Итого").font = Font(name=FONT, bold=True)
            for c in range(1, ncols):
                fmt = col_formats[c]
                summable = (
                    bool(fmt) and "%" not in fmt and "DD" not in fmt
                )  # проценты и даты не суммируем
                numeric_or_formula = any(
                    isinstance(v, int | float) or (isinstance(v, str) and v.startswith("="))
                    for v, _ in (row[c] for row in parsed)
                )
                if summable and numeric_or_formula:
                    letter = get_column_letter(c + 1)
                    cell = ws.cell(
                        row=total_row,
                        column=c + 1,
                        value=f"=SUM({letter}{first_data}:{letter}{last_data})",
                    )
                    cell.font = Font(name=FONT, bold=True)
                    cell.number_format = col_formats[c] or "General"
                    cell.alignment = Alignment(horizontal="right")
            for c in range(1, ncols + 1):
                ws.cell(row=total_row, column=c).border = Border(
                    top=Side(style="medium", color=palette.primary)
                )

        # ширина колонок по содержимому (с потолком)
        for c in range(ncols):
            texts = ([headers[c]] if c < len(headers) else []) + [
                str(r[c]) for r in rows[:300] if c < len(r)
            ]
            width = min(60, max(10, max((len(t) for t in texts), default=8) * 1.15 + 2))
            ws.column_dimensions[get_column_letter(c + 1)].width = width
        if headers and rows:
            ws.auto_filter.ref = f"A1:{get_column_letter(ncols)}{last_data}"
        if isinstance(sheet_data.get("chart"), dict) and rows and headers:
            _chart(ws, sheet_data["chart"], header_row, first_data, last_data, ncols)

    if not wb.worksheets:
        wb.create_sheet("Лист")
    stream = io.BytesIO()
    wb.save(stream)
    return stream.getvalue()


__all__ = ["SPREADSHEET_TOOL", "build_spreadsheet", "is_safe_formula", "parse_cell"]
