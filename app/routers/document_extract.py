"""Извлечение текста из загруженного документа для подстановки в промпт чата.

Намеренно ИЗОЛИРОВАННЫЙ роут — не трогает /v1/chat, классификатор,
YAML-роутер или ChatRequest/ChatResponse. Фронтенд сначала вызывает
POST /v1/documents/extract, получает обратно текст, сам приклеивает его
к началу prompt пользователя и уже этот составной prompt отправляет
обычным вызовом POST /v1/chat. Контракт существующего чата не меняется.

Ничего не пишет на диск и в БД — файл живёт только в памяти на время
запроса (в отличие от app/routers/parser.py в ERP_Bakend, которому
нужна персистентность для Celery-обработки — здесь она не нужна).

Библиотеки — те же, что уже используются в
app/services/procurement_parser/ (python-docx, pdfplumber, openpyxl),
дополнительных зависимостей не требуется.
"""
from __future__ import annotations

import io

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile, status

from app.auth import CurrentUser, get_current_user

router = APIRouter(prefix="/v1/documents", tags=["documents"])

MAX_FILE_SIZE_BYTES = 15 * 1024 * 1024  # 15 МБ — с запасом для PDF со сканами
MAX_EXTRACTED_CHARS = 20_000  # чтобы не взорвать бюджет токенов промпта

SUPPORTED_EXTENSIONS = {".docx", ".pdf", ".xlsx"}


def _extract_docx(content: bytes) -> str:
    import docx  # python-docx; lazy import, как в word_extractor.py

    document = docx.Document(io.BytesIO(content))
    parts = [p.text for p in document.paragraphs if p.text.strip()]

    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))

    return "\n".join(parts)


def _extract_pdf(content: bytes) -> str:
    from app.utils.pdf import extract_pdf_pages

    return "\n".join(extract_pdf_pages(content))


def _extract_xlsx(content: bytes) -> str:
    from openpyxl import load_workbook

    workbook = load_workbook(filename=io.BytesIO(content), read_only=True, data_only=True)
    parts: list[str] = []

    for sheet in workbook.worksheets:
        parts.append(f"# Лист: {sheet.title}")
        for row in sheet.iter_rows(values_only=True):
            cells = [str(value) for value in row if value is not None]
            if cells:
                parts.append(" | ".join(cells))

    return "\n".join(parts)


EXTRACTORS = {
    ".docx": _extract_docx,
    ".pdf": _extract_pdf,
    ".xlsx": _extract_xlsx,
}


@router.post("/extract")
async def extract_document_text(
    file: UploadFile = File(...),
    user: CurrentUser = Depends(get_current_user),
) -> dict:
    """Извлечь текст из документа. Ничего не сохраняет — только текст в ответе."""
    if not file.filename:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, detail="Имя файла отсутствует")

    suffix = "." + file.filename.rsplit(".", 1)[-1].lower() if "." in file.filename else ""
    if suffix not in SUPPORTED_EXTENSIONS:
        raise HTTPException(
            status.HTTP_400_BAD_REQUEST,
            detail=f"Поддерживаются только файлы: {', '.join(sorted(SUPPORTED_EXTENSIONS))}",
        )

    content = await file.read()
    if len(content) > MAX_FILE_SIZE_BYTES:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, detail="Файл слишком большой (лимит 15 МБ)")

    try:
        text = EXTRACTORS[suffix](content)
    except Exception as exc:
        raise HTTPException(
            status.HTTP_400_BAD_REQUEST,
            detail=f"Не удалось прочитать файл: {exc}",
        ) from exc

    truncated = len(text) > MAX_EXTRACTED_CHARS
    if truncated:
        text = text[:MAX_EXTRACTED_CHARS]

    return {
        "filename": file.filename,
        "text": text,
        "truncated": truncated,
        "char_count": len(text),
    }
